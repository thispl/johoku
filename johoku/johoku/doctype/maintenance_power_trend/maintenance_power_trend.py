# Copyright (c) 2024, TEAMPRO and contributors
# For license information, please see license.txt

# import frappe
from frappe.model.document import Document
import frappe
import openpyxl
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side
from openpyxl.styles import GradientFill, PatternFill
from six import BytesIO
import io  # Add this import to resolve the NameError

from datetime import datetime
from datetime import datetime, timedelta
from frappe.utils import (getdate, cint, add_months, date_diff, add_days, format_date,
    nowdate, get_datetime_str, cstr, get_datetime, now_datetime, format_datetime)
from openpyxl.drawing.image import Image as xlImage
import requests

class MaintenancePowerTrend(Document):
    pass

@frappe.whitelist()
def download():
    filename = 'Maintenance Power Trend.xlsx'
    build_xlsx_response(filename)

def make_xlsx(sheet_name=None):
    args = frappe.local.form_dict
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_name if sheet_name else 'Sheet1'
    
    data = get_data(args)
    
    if not data:  # Check if data is empty or None
        frappe.throw("No data available to generate the report.")
    
    
    image_url="http://157.245.101.198/files/JohokuNew.png"
    image_data = requests.get(image_url).content
    img = xlImage(io.BytesIO(image_data))
    img_cell = ws.cell(row=1, column=7)
    img.width = 100
    img.height = 80
    ws.add_image(img,"I1")
         
        
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=10)
    ws["K1"].value = "JOHOKU MANUFACTURING PRIVATE LIMITED"
    ws["K1"].alignment = Alignment(horizontal='left', vertical='center', wrap_text=True)
    ws["K1"].font = Font(bold=True, size=14)
    ws.merge_cells(start_row=1, start_column=11, end_row=1, end_column=24)
    for row in data:
        if not isinstance(row, (list, tuple)):
            row = [row]  
        ws.append(row)

    
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=24)
    ws.merge_cells(start_row=3, start_column=1, end_row=4, end_column=1)
    ws.merge_cells(start_row=3, start_column=2, end_row=3, end_column=6)
    ws.merge_cells(start_row=3, start_column=7, end_row=3, end_column=11)
    ws.merge_cells(start_row=3, start_column=12, end_row=3, end_column=16)
    # ws.merge_cells(start_row=3, start_column=x, end_row=4, end_column=x)
    for x in range(17,25):
        ws.merge_cells(start_row=3, start_column=x, end_row=4, end_column=x)
    for x in range(2,18):
        ws.merge_cells(start_row=4, start_column=x, end_row=4, end_column=x)
    apply_styles(ws)
    set_column_widths(ws)

    xlsx_file = BytesIO()
    wb.save(xlsx_file)
    xlsx_file.seek(0)
    return xlsx_file

def apply_styles(ws):
    align_center = Alignment(horizontal='center', vertical='center', wrap_text=True)
    align_right = Alignment(horizontal='right', vertical='top', wrap_text=True)
    align_left = Alignment(horizontal='left', vertical='top', wrap_text=True)
    header_font = Font(bold=True, size=14)
    text_font = Font(bold=True, size=8)
    text_font_header = Font(bold=True, size=10)
    text_font_header_data = Font(bold=True, size=9)
    text_font_data = Font(bold=False, size=9)
    thin_border = Side(border_style="thin", color="000000")
    border = Border(
        left=Side(border_style='thin'),
        right=Side(border_style='thin'),
        top=Side(border_style='thin'),
        bottom=Side(border_style='thin')
    )
    outline_border = Border(
    left=thin_border,
    right=thin_border,
    top=thin_border,
    bottom=thin_border
)
    outline_border_new = Border(
    left=thin_border,
    top=thin_border,
)

# Apply the outline border for the entire range
    outline_range = ws.iter_rows(min_row=1, max_row=1, min_col=24, max_col=24)
    for rows in outline_range:
        for cell in rows:
            cell.border = outline_border_new
    outline_range = ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=1, max_col=24)
    for rows in outline_range:
        for cell in rows:
            cell.border = outline_border
    for rows in ws.iter_rows(min_row=2, max_row=2, min_col=1, max_col=24):
        for cell in rows:
            cell.font = header_font
            cell.alignment = align_center
            cell.border = border
    for x in range(3,5):
        for rows in ws.iter_rows(min_row=x, max_row=x, min_col=1, max_col=24):
            for cell in rows:
                cell.font = text_font_header
                cell.alignment = align_center
                cell.border = border
    for x in range(5,ws.max_row + 1):
        for rows in ws.iter_rows(min_row=x, max_row=x, min_col=1, max_col=24):
            for cell in rows:
                cell.font = text_font_data
                cell.alignment = align_center
                cell.border = border
    # for rows in ws.iter_rows(min_row=3, max_row=4, min_col=1, max_col=25):
    #     for cell in rows:
    #         cell.fill = PatternFill(fgColor="f4f087", fill_type = "solid")
    
    

def set_column_widths(ws):
    column_widths = [11] *24
    for i, width in enumerate(column_widths, start=1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = width
    row_heights = [60, 60, 30,30,30] + [20] * (ws.max_row - 1) +[30]*1 
    for i, height in enumerate(row_heights, start=1):
        ws.row_dimensions[i].height = height

def build_xlsx_response(filename):
    xlsx_file = make_xlsx(sheet_name=filename)
    frappe.response['filename'] = filename + '.xlsx'
    frappe.response['filecontent'] = xlsx_file.getvalue()
    frappe.response['type'] = 'binary'


# def get_data(args):
#     from_date = args.get('from_date')
#     to_date = args.get('to_date')
def get_data():
    from_date ="2024-11-01"
    to_date ="2024-11-30"
    # from_date ="2024-01-01"
    # to_date ="2024-12-31"
    
    

    if not from_date or not to_date:
        frappe.throw("From Date and To Date are required.")
    # from_date_str =from_date
    data = [
            ["Maintenance Power Trend"],
            ["Month",'STAFF',"","","","",'TECH',"","","","", 'TRAINEE',"","","","","Total W1","Total W2","Total W3","Total W4","Total W5",
            "Updated By","Approved By","Remarks"],
            ["","W1","W2","W3","W4","W5","W1","W2","W3","W4","W5","W1","W2","W3","W4","W5"]
        ]
    # month = int(from_date.month)
    # year = int(from_date.year)
    # from_date = from_date.date()
    # to_date = to_date.date()
    
    month_list = get_months_between_dates(from_date, to_date)
    month_list = get_months_between_dates(from_date, to_date)
    month_year_list = get_month_with_year_between_dates(from_date, to_date)
    count = 0
    index = 0
    for month_year in month_year_list:
        print(month_year)
        count += 1        
        week_list = get_week_count_from_month(month_year)
        row = []
        plan_total_staff_row = []
        plan_total_tech_row = []
        plan_total_trainee_row = []
        row1 =[]
        row2 =[]
        row3 =[]
        week1_row =[]
        week2_row =[]
        week3_row = []
        week4_row =[]
        week5_row =[]
        total_week_row =[]
        month_row =[]
        week1 = week2 = week3 = week4 = week5 = 0
        
        
        # for month_data in month_list:
        month_list_data = month_list[index]
        month_row.append(month_list_data)
        print(month_row)
        
        week_number = 0
        for week_date in week_list:
            print(week_date)
            from_date = week_date.get('start_date')
            to_date = week_date.get('end_date')
            week_number_ = week_date.get('week_number')
            week_number += 1
            plan_total_staff = frappe.db.count("Shift Assignment", {
                'start_date': ['between',[from_date,to_date]],
                'docstatus': 1,
                'employee_category': "Staff"
            })
            plan_total_tech = frappe.db.count("Shift Assignment", {
                'start_date': ['between',[from_date,to_date]],
                'docstatus': 1,
                'employee_category': "ITI"
            })
            plan_total_tt = frappe.db.count("Shift Assignment", {
                'start_date': ['between',[from_date,to_date]],
                'docstatus': 1,
                'employee_category': "TT"
            })
            plan_total_gt = frappe.db.count("Shift Assignment", {
                'start_date': ['between',[from_date,to_date]],
                'docstatus': 1,
                'employee_category': "GT"
            })
            plan_total_trainee = plan_total_tt + plan_total_gt
            print(plan_total_staff)
            print(plan_total_tech)
            print(plan_total_trainee)
            if week_number == 1:
                print(week1)
                week1 += plan_total_staff + plan_total_tech + plan_total_trainee
            elif week_number == 2:
                week2 += plan_total_staff + plan_total_tech + plan_total_trainee
            elif week_number == 3:
                week3 += plan_total_staff + plan_total_tech + plan_total_trainee
            elif week_number == 4:
                week4 += plan_total_staff + plan_total_tech + plan_total_trainee
            elif week_number == 5:
                week5 += plan_total_staff + plan_total_tech + plan_total_trainee
            print(week_number)
            print(week1)
            print(week2)
            plan_total_staff_row.append(plan_total_staff)
            plan_total_trainee_row.append(plan_total_trainee)
            plan_total_tech_row.append(plan_total_tech)
            print(plan_total_staff_row)
        index += 1
        # print(plan_total_staff_row)
        week1_row.append(week1)
        week2_row.append(week2)
        week3_row.append(week3)
        week4_row.append(week4)
        week5_row.append(week5)
        # print(week1_row)
        total_week_row = week1_row + week2_row + week3_row + week4_row + week5_row
        # print(total_week_row)
        row1 += plan_total_staff_row
        row2 += plan_total_tech_row
        row3 += plan_total_trainee_row
        row = month_row + row1 + row2 + row3 + total_week_row
        data.append(row)
    # print(count)       
    week_row = []
    # return data


@frappe.whitelist()
def get_dates(from_date ,to_date):
    no_of_days = date_diff(add_days(to_date, 1), from_date)
    dates = [add_days(from_date, i) for i in range(0, no_of_days)]
    return dates


import frappe
import calendar
# from frappe.utils import getdate, add_months

@frappe.whitelist()
def get_months_between_dates(from_date, to_date):
    from_date = getdate(from_date)
    to_date = getdate(to_date)

    months = []
    current_date = from_date

    while current_date <= to_date:
        month_name = calendar.month_abbr[current_date.month]
        if month_name not in months:
            months.append(month_name)
        current_date = add_months(current_date, 1)

    return months

from datetime import datetime

from datetime import datetime, timedelta
import calendar

def get_week_count_from_month(month_year):
    # Parse "YYYY-MM" into a datetime object
    first_day_of_month = datetime.strptime(month_year, '%Y-%m')
    
    # Get the last day of the month
    _, last_day = calendar.monthrange(first_day_of_month.year, first_day_of_month.month)
    last_day_of_month = first_day_of_month.replace(day=last_day)

    # Calculate week count and their start/end dates
    weeks = []
    current_day = first_day_of_month

    # Align the current_day to the first Monday of the month
    while current_day.weekday() != 0:  # 0 = Monday
        current_day += timedelta(days=1)

    while current_day <= last_day_of_month:
        # Start date of the week is the current day
        week_start = current_day

        # End date of the week (Saturday)
        week_end = current_day + timedelta(days=5)

        # Ensure the week_end doesn't exceed the month's last day
        if week_end > last_day_of_month:
            week_end = last_day_of_month

        # Add the week to the list
        weeks.append({
            'week_number': week_start.isocalendar()[1],  # ISO week number
            'start_date': week_start.strftime('%Y-%m-%d'),
            'end_date': week_end.strftime('%Y-%m-%d')
        })

        # Move to the next week (next Monday)
        current_day = week_end + timedelta(days=2)

    return weeks


@frappe.whitelist()
def get_month_with_year_between_dates(from_date, to_date):
    from_date = getdate(from_date)
    to_date = getdate(to_date)

    months = []
    current_date = from_date

    while current_date <= to_date:
        month_year = current_date.strftime("%Y-%m")  # Format: "YYYY-MM"
        if month_year not in months:
            months.append(month_year)
        current_date = add_months(current_date, 1)

    return months

