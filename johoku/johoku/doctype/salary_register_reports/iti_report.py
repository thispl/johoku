# Copyright (c) 2024, TEAMPRO and contributors
# For license information, please see license.txt


from frappe.model.document import Document
import frappe
import openpyxl
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side
from openpyxl.styles import GradientFill, PatternFill
from six import BytesIO
import io
from datetime import datetime
from datetime import datetime, timedelta
from frappe.utils import (getdate, cint, add_months, date_diff, add_days, format_date,
    nowdate, get_datetime_str, cstr, get_datetime, now_datetime, format_datetime)
from openpyxl.drawing.image import Image as xlImage
import requests


@frappe.whitelist()
def download():
    filename = 'Salary Register for ITI.xlsx'
    build_xlsx_response(filename)

def make_xlsx(sheet_name=None):
    args = frappe.local.form_dict
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_name if sheet_name else 'Sheet1'
    
    data = get_data(args)
    
    if not data:
        frappe.throw("No data available to generate the report.")

    image_url="http://157.245.101.198/files/JohokuNew.png"
    # image_data = requests.get(image_url).content
    image_data = requests.get(image_url, verify=False).content
    img = xlImage(io.BytesIO(image_data))
    img_cell = ws.cell(row=1, column=10)
    img.width = 100
    img.height = 80
    ws.add_image(img,"J1")
         
        
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=11)
    ws["L1"].value = "JOHOKU MANUFACTURING PRIVATE LIMITED"
    ws["L1"].alignment = Alignment(vertical='center', wrap_text=True)
    ws["L1"].font = Font(bold=True, size=14)
    ws["AT1"].value = "Prepared By"
    ws["AX1"].value = "Checked By"
    ws["BE1"].value = "Verified By"
    ws["BJ1"].value = "Approved By"
    
    for row in data:
        if not isinstance(row, (list, tuple)):
            row = [row]
        ws.append(row)

    ws.merge_cells(start_row=1, start_column=12, end_row=1, end_column=45)
    # ws.merge_cells(start_row=1, start_column=41, end_row=1, end_column=44)
    ws.merge_cells(start_row=1, start_column=46, end_row=1, end_column=49)
    ws.merge_cells(start_row=1, start_column=50, end_row=1, end_column=56)
    ws.merge_cells(start_row=1, start_column=57, end_row=1, end_column=61)
    ws.merge_cells(start_row=1, start_column=62, end_row=1, end_column=66)
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=45)
    # ws.merge_cells(start_row=2, start_column=41, end_row=3, end_column=44)
    ws.merge_cells(start_row=2, start_column=46, end_row=3, end_column=49)
    ws.merge_cells(start_row=2, start_column=50, end_row=3, end_column=52)
    ws.merge_cells(start_row=2, start_column=53, end_row=3, end_column=56)
    ws.merge_cells(start_row=2, start_column=57, end_row=3, end_column=61)
    ws.merge_cells(start_row=2, start_column=62, end_row=3, end_column=66)
    ws.merge_cells(start_row=3, start_column=1, end_row=3, end_column=45)
    # ws.merge_cells(start_row=3, start_column=40, end_row=3, end_column=43)
    for x in range(1,17):
        ws.merge_cells(start_row=4, start_column=x, end_row=5, end_column=x)
    for x in range(27,37):
        ws.merge_cells(start_row=4, start_column=x, end_row=5, end_column=x)
    ws.merge_cells(start_row=4, start_column=17, end_row=4, end_column=26)
    ws.merge_cells(start_row=4, start_column=37, end_row=4, end_column=46)
    ws.merge_cells(start_row=4, start_column=47, end_row=4, end_column=56)
    ws.merge_cells(start_row=4, start_column=58, end_row=4, end_column=64)

    
    # for x in range(1,17):
    #     ws.merge_cells(start_row=4, start_column=x, end_row=5, end_column=x)
    # for x in range(27,37):
    #     ws.merge_cells(start_row=4, start_column=x, end_row=5, end_column=x)
    
    # ws.merge_cells(start_row=4, start_column=17, end_row=4, end_column=26)
    # ws.merge_cells(start_row=4, start_column=37, end_row=4, end_column=46)
    # ws.merge_cells(start_row=4, start_column=47, end_row=4, end_column=56)
    # ws.merge_cells(start_row=4, start_column=58, end_row=4, end_column=64)
    # ws.merge_cells(start_row=4, start_column=50, end_row=5, end_column=50)
    ws.merge_cells(start_row=4, start_column=57, end_row=5, end_column=57)
    ws.merge_cells(start_row=4, start_column=65, end_row=5, end_column=65)
    ws.merge_cells(start_row=4, start_column=66, end_row=5, end_column=66)

    apply_styles(ws,args)
    set_column_widths(ws)

    xlsx_file = BytesIO()
    wb.save(xlsx_file)
    xlsx_file.seek(0)
    return xlsx_file

def apply_styles(ws,args):
    from_date = args.get('from_date')
    to_date = args.get('to_date')

    align_center = Alignment(horizontal='center', vertical='center', wrap_text=True)
    # align_right = Alignment(horizontal='right', vertical='top', wrap_text=True)
    # align_left = Alignment(horizontal='left', vertical='top', wrap_text=True)
    header_font = Font(bold=True, size=30)
    header_font_1 = Font(bold=True, size=20)
    # text_font = Font(bold=True, size=8)
    text_font_header = Font(bold=True, size=10)
    # text_font_data = Font(bold=False, size=9)
    border = Border(
        left=Side(border_style='thin'),
        right=Side(border_style='thin'),
        top=Side(border_style='thin'),
        bottom=Side(border_style='thin')
    )
    salary_slips = frappe.db.sql("""
        SELECT * 
        FROM `tabSalary Slip`
        WHERE start_date = %s AND end_date = %s AND docstatus != 2 AND salary_structure = 'DT-ITI Structure'
    """, (from_date, to_date), as_dict=True)
    index=0
    for s in salary_slips:
        if not isinstance(s, dict):
            continue

        emp_data = frappe.get_list(
            'Employee',
            filters={'name': s['employee'],'employee_category' : "ITI"},
            fields=["date_of_birth", 'date_of_joining', 'bank_ac_no', 'salary_mode', 'employee_category',
            'basic','house_rent_allowance','conveyance_allowance','education_allowance','food_allowance',
            'medical_allowance','special_allowance','fixed_earning'
            ],
            limit_page_length=1
        )

        emp = emp_data[0] if emp_data else frappe._dict({
            "date_of_birth": None,
            "date_of_joining": None,
            "bank_ac_no": 0,
            "salary_mode": 0,
            "employee_category": 0,
            "basic":0,
            "house_rent_allowance":0,
            "conveyance_allowance":0,
            "education_allowance":0,
            "food_allowance":0,
            "medical_allowance":0,
            "special_allowance":0,
            "fixed_earning":0,
        })
        if emp.employee_category == "ITI":
            index+=1
            
    for x in range(1,4):
        for rows in ws.iter_rows(min_row=x, max_row=x, min_col=1, max_col=66):
            for cell in rows:
                cell.font = header_font
                cell.alignment = align_center
                cell.border = border
    for x in range(3, 4):
        for rows in ws.iter_rows(min_row=x, max_row=x, min_col=1, max_col=57):
            for cell in rows:
                cell.font = header_font_1
                cell.alignment = align_center
                cell.border = border
    for rows in ws.iter_rows(min_row=4, max_row=5, min_col=1, max_col=15):
        for cell in rows:
            cell.font = text_font_header
            cell.alignment = align_center
            cell.border = border
    for rows in ws.iter_rows(min_row=4, max_row=5, min_col=25, max_col=36):
        for cell in rows:
            cell.font = text_font_header
            cell.alignment = align_center
            cell.border = border
    for rows in ws.iter_rows(min_row=4, max_row=5, min_col=49, max_col=49):
        for cell in rows:
            cell.font = text_font_header
            cell.alignment = align_center
            cell.border = border
    for rows in ws.iter_rows(min_row=4, max_row=5, min_col=57, max_col=57):
        for cell in rows:
            cell.font = text_font_header
            cell.alignment = align_center
            cell.border = border
    for rows in ws.iter_rows(min_row=4, max_row=4, min_col=16, max_col=25):
        for cell in rows:
            cell.font = text_font_header
            cell.alignment = align_center
            cell.border = border
    for rows in ws.iter_rows(min_row=4, max_row=4, min_col=36, max_col=47):
        for cell in rows:
            cell.font = text_font_header
            cell.alignment = align_center
            cell.border = border
    for rows in ws.iter_rows(min_row=4, max_row=4, min_col=50, max_col=66):
        for cell in rows:
            cell.font = text_font_header
            cell.alignment = align_center
            cell.border = border
    for x in range(16,26):
        for rows in ws.iter_rows(min_row=5, max_row=5, min_col=x, max_col=x):
            for cell in rows:
                cell.font = text_font_header
                cell.alignment = align_center
                cell.border = border
    for x in range(36,47):
        for rows in ws.iter_rows(min_row=5, max_row=5, min_col=x, max_col=x):
            for cell in rows:
                cell.font = text_font_header
                cell.alignment = align_center
                cell.border = border
    for x in range(47,67):
        for rows in ws.iter_rows(min_row=5, max_row=5, min_col=x, max_col=x):
            for cell in rows:
                cell.font = text_font_header
                cell.alignment = align_center
                cell.border = border

    for rows in ws.iter_rows(min_row=5, max_row=5, min_col=48, max_col=48):
        for cell in rows:
            cell.font = text_font_header
            cell.alignment = align_center
            cell.border = border

    for rows in ws.iter_rows(min_row=5, max_row=5, min_col=56, max_col=56):
        for cell in rows:
            cell.font = text_font_header
            cell.alignment = align_center
            cell.border = border

    for rows in ws.iter_rows(min_row=6, max_row=ws.max_row, min_col=1, max_col=66):
        for cell in rows:
            cell.border = border
            cell.alignment = align_center

    ws.merge_cells(start_row=6+index, start_column=1, end_row=6+index, end_column=15)
    for rows in ws.iter_rows(min_row=4, max_row=5, min_col=1, max_col=66):
        for cell in rows:
            cell.fill = PatternFill(fgColor="f4f087", fill_type = "solid")
    for rows in ws.iter_rows(min_row=6, max_row=5+index, min_col=16, max_col=16):
        for cell in rows:
            cell.fill = PatternFill(fgColor="FABF8F", fill_type = "solid")
    for rows in ws.iter_rows(min_row=6, max_row=5+index, min_col=27, max_col=35):
        for cell in rows:
            cell.fill = PatternFill(fgColor="FABF8F", fill_type = "solid")
    for rows in ws.iter_rows(min_row=6, max_row=5+index, min_col=47, max_col=53):
        for cell in rows:
            cell.fill = PatternFill(fgColor="FABF8F", fill_type = "solid")
    for rows in ws.iter_rows(min_row=6, max_row=5+index, min_col=55, max_col=55):
        for cell in rows:
            cell.fill = PatternFill(fgColor="FABF8F", fill_type = "solid")
    for rows in ws.iter_rows(min_row=6, max_row=5+index, min_col=64, max_col=64):
        for cell in rows:
            cell.fill = PatternFill(fgColor="FABF8F", fill_type = "solid")
    for rows in ws.iter_rows(min_row=6+index, max_row=6+index, min_col=1, max_col=66):
        for cell in rows:
            cell.fill = PatternFill(fgColor="f4f087", fill_type = "solid")
    
    

def set_column_widths(ws):
    column_widths = [8] *1 + [12] * 4 + [15] * 2 + [15]  + [12]* 56 + [14]* 10
    # column_widths = [6, 5] + [5] * 13 + [7] * 2 + [5] * 1 +[6] * 1 + [5] * 4 +[6] * 1 + [4] * 5+ [8] * 7
    for i, width in enumerate(column_widths, start=1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = width
    max_row_count = ws.max_row  - 6
    row_heights = [80, 80, 80, 25, 35] + [30] * max_row_count +[30]*1 # Example row heights
    for i, height in enumerate(row_heights, start=1):
        ws.row_dimensions[i].height = height

def build_xlsx_response(filename):
    xlsx_file = make_xlsx(sheet_name=filename)
    frappe.response['filename'] = filename + '.xlsx'
    frappe.response['filecontent'] = xlsx_file.getvalue()
    frappe.response['type'] = 'binary'

def get_leave_count(employee, leave_type, from_date, to_date):
    if not leave_type:
        frappe.throw("Leave Type is required.")

    # Ensure date format consistency (use getdate if necessary)
    from_date = getdate(from_date)
    to_date = getdate(to_date)
    full_day_leave_count = halfdays_leave_count = leave_count = 0
    full_day_leave_count = frappe.db.count("Attendance", {
        'employee': employee,
        'leave_type': leave_type,
        'status': 'On Leave',
        'leave_application':['!=',''],
        'attendance_date': ['between', [from_date, to_date]],
        'docstatus':('!=','2')
    })
    halfdays_leave_count = frappe.db.count("Attendance", {
        'employee': employee,
        'leave_type': leave_type,
        'status': 'Half Day',
        'leave_application':['!=',''],
        'attendance_date': ['between', [from_date, to_date]],
        'docstatus':('!=','2')
    })
    leave_count = full_day_leave_count + (halfdays_leave_count / 2)
    return leave_count


def get_holiday_count(employee, from_date, to_date):
    holiday_list = frappe.db.get_value('Employee', employee, 'holiday_list')
    total_weekoff = 0
    total_public_holiday = 0
    total_other_holiday = 0

    if holiday_list:
        holiday_records = frappe.db.sql("""
            SELECT holiday_date, weekly_off, holiday_type
            FROM `tabHoliday`
            WHERE parent = %s AND holiday_date BETWEEN %s AND %s
        """, (holiday_list, from_date, to_date), as_dict=True)

        for record in holiday_records:
            if record['weekly_off'] == 1:
                total_weekoff += 1
            elif record['holiday_type'] == "PH":
                total_public_holiday += 1
            else:
                total_other_holiday += 1

    return {
        'total_weekoff': total_weekoff,
        'total_public_holiday': total_public_holiday,
        'total_other_holiday': total_other_holiday,
        'total_holiday': total_weekoff + total_public_holiday
    }
        
  
def get_data(args):
    from_date = args.get('from_date')
    to_date = args.get('to_date')

    if not from_date or not to_date:
        frappe.throw("From Date and To Date are required.")

    to_date_str =to_date

    if isinstance(to_date_str, str):
        month_year_format_str = datetime.strptime(to_date_str, '%Y-%m-%d')
    else:
        month_year_format_str = to_date_str

    month_year_format = month_year_format_str.strftime('%B %Y')
   
    
    data = [
            [""],

            ["SALARY STATEMENT FOR THE MONTH OF: " + month_year_format.upper() + " - ITI"],

            ["S.No", "Emp Code", "Name", "DOB", "Grade", "DOJ", "Department", "Designation", "Aadhar No", "PAN No", "UAN No", "ESIC No.", "Bank A/C No", 
             "Bank Type","Category", "No Working day", "Fixed Gross"] + 
            [""] * 9 + ["Present Days", "S/PH", "COFF", "CL", "SL", "EL", "N/F", "LOP", "Absent", "NPD", 
             "GROSS EARNINGS"] + [""] * 9 + ["VARIABLE EARNINGS"] + [""] * 9 + 
            ["Total Earnings", "Deductions"] + [""] * 6 +["Total Deductions", "Net Pay INR"],

            [""] * 16 + 
            ["Basic", "HRA", "Conveyance", "Education", "Food Allowance", "Attire Allowance","Mobile Allowance","Medical Reimbursement",
             "Spl Allow", "Total Fixed Gross Earning"] + 
            [""] * 10 + ["Basic", "HRA", "Conveyance", "Education", "Food Allowance", "Attire Allowance","Mobile Allowance",
             "Medical Reimbursement", "Spl Allow", "Total Gross Earning", 
             "Attendance Bonus", "Heat Allowance", "TL Allowance", "EL Encashment", "Gratuity Pay","Travel Allowance", "Bonus", "OT Hrs", 
             "OT Amount","Arear Allowance", "", "EPF Round Off", "ESIC Round up off", "PT",  
             "LWF", "Loan", "Other Deductions", "TDS Deduction"]
        ]

    salary_slips = frappe.db.sql("""
    SELECT * 
    FROM `tabSalary Slip`
    WHERE start_date = %s AND end_date = %s AND docstatus != 2 AND salary_structure = 'DT-ITI Structure'
""", (from_date, to_date), as_dict=True)

    earned_gross = total_earnings = total = net_pay = 0
    working_days_total = fixed_basic_total = fixed_hra_total = fixed_conveyance_total = fixed_education_total = fixed_food_total = fixed_attire_total = fixed_mobile_total = fixed_medical_total = 0
    fixed_special_total = fixed_gross_total = present_days_total = sph_total = coff_total = cl_total = sl_total = el_total = nf_total = 0
    lop_total = npd_total = gross_basic_total = gross_hra_total = gross_conveyance_total = gross_education_total = gross_food_total = gross_attire_total = gross_mobile_total = gross_medical_total = 0
    gross_special_total = earned_gross_total = attendance_allow_total = heat_allow_total = tl_allow_total = gratuity_total = other_allow_total = ot_hours = overtime_total =travel_allow= total_earnings_total = epf_total = 0
    prof_tax_total = esci_total = loan_total = lwf_total = tds_total = others_total = total_total = net_pay_total = absent_total = 0
    index = epf_ac_mo = 0
    el_enhancement = 0
    total_arrear_allowance =0
    row_total = []
    for s in salary_slips:
        if not isinstance(s, dict):
            continue

        emp_data = frappe.get_list(
            'Employee',
            filters={'name': s['employee'],'employee_category' : "ITI"},
            fields=["date_of_birth", 'date_of_joining', 'bank_ac_no', 'salary_mode', 'employee_category',
            'basic','house_rent_allowance','conveyance_allowance','education_allowance','food_allowance',
            'medical_allowance','special_allowance','fixed_earning', 'grade', 'bank_type' , 'esic_no','pan_no', "uan_no", "aadhar_no"
            ],
            limit_page_length=1
        )

        emp = emp_data[0] if emp_data else frappe._dict({
            "date_of_birth": None,
            "date_of_joining": None,
            "bank_ac_no": 0,
            "salary_mode": 0,
            "employee_category": 0,
            "basic":0,
            "house_rent_allowance":0,
            "conveyance_allowance":0,
            "education_allowance":0,
            "food_allowance":0,
            "medical_allowance":0,
            "special_allowance":0,
            "fixed_earning":0,
        })

        holiday_data = get_holiday_count(s['employee'], from_date, to_date)
        present_doc_name_list = frappe.db.get_all(
            "Attendance", 
            filters={
                'employee': s['employee'],
                'status': 'Present',  
                'attendance_date': ["between", [from_date, to_date]],
                'docstatus': ["!=", 2], 
            },
            fields=['attendance_date', 'name', 'employee']
        )

        present_days_count_ = 0
        halfdays_count = 0
        halfdays_absent_count = 0
        for doc_name in present_doc_name_list:
            attendance_date = doc_name['attendance_date']
            if not check_holiday(attendance_date, s['employee']):
                present_doc_name = frappe.db.get_value(
                    "Attendance", 
                    filters={
                        'employee': s['employee'],
                        'status': ['in',['Present','Half Day']],   
                        'attendance_date': attendance_date,
                        'docstatus': ["!=", 2], 
                    },
                    fieldname=['attendance_date', 'name']
                )
                if present_doc_name:
                    present_days_count_ += 1
                halfday_doc_name = frappe.db.get_value(
                    "Attendance", 
                    filters={
                        'employee': s['employee'],
                        'status': 'Half Day',  
                        'attendance_date': attendance_date,
                        'docstatus': ["!=", 2], 
                        'total_working_hours':[">=",4]
                    },
                    fieldname=['attendance_date', 'name']
                )
                if halfday_doc_name:
                    halfdays_count += 1
                halfday_absent_doc_name = frappe.db.get_value(
                    "Attendance", 
                    filters={
                        'employee': s['employee'],
                        'status': 'Half Day',  
                        'attendance_date': attendance_date,
                        'docstatus': ["!=", 2], 
                        'total_working_hours':["<",4]
                    },
                    fieldname=['attendance_date', 'name']
                )
                if halfday_absent_doc_name:
                    halfdays_absent_count += 1

        present_days = present_days_count_ + (halfdays_count / 2)
        coff_count = get_leave_count(s['employee'], "Compensatory Off", from_date, to_date)
        sl_count = get_leave_count(s['employee'], "Sick Leave", from_date, to_date)
        el_count = get_leave_count(s['employee'], "Earned Leave", from_date, to_date)
        cl_count = get_leave_count(s['employee'], "Casual Leave", from_date, to_date)
        lop_count = get_leave_count(s['employee'], "Leave Without Pay", from_date, to_date)
        # lop_count = (halfdays_absent_count / 2) + lop_count_leave
        components = ['Basic', 'House Rent Allowance', 'Conveyance', 'Education Allowance',
            'Food Allowance', 'Medical Re imbursement', 'Special Allowance',
            'Mobile Allowance', 'Heat Allowance', 'Attire Allowance', 'Mobile Allowance', 'TL Allowance', 'TDS Deduction', 'Gratuity','Arear Allowance',
            'Overtime','Travel Allowance', 'Attendance Allowance', 'Provident Fund','Arear Allowance',
                    'Employer State Insurance', 'Professional Tax', 'Others', 'Labour Welfare Fund', 'Loan','EL Encashment']

        salary_details = {}
        for comp in components:
            salary_amount = frappe.get_value("Salary Detail", {'parent': s['name'], 'salary_component': comp}, 'amount') or 0.0
            salary_details[comp] = salary_amount
         
        earned_gross = sum([
            salary_details.get('Basic', 0.0) + salary_details.get('House Rent Allowance', 0.0) + salary_details.get('Conveyance', 0.0) + salary_details.get('Education Allowance', 0.0) +
            salary_details.get('Food Allowance', 0.0) + salary_details.get('Attire Allowance', 0.0) + salary_details.get('Mobile Allowance', 0.0) + salary_details.get('Medical Re imbursement', 0.0) + salary_details.get('Special Allowance', 0.0)
        ])

        total_earnings = sum([
            salary_details.get('Attendance Allowance', 0.0) +
            salary_details.get('Heat Allowance', 0.0) +
             salary_details.get('Gratuity', 0.0) +
            salary_details.get('TL Allowance', 0.0) +
            salary_details.get('Overtime', 0.0) +
            salary_details.get('Travel Allowance', 0.0) +
            salary_details.get('EL Encashment', 0.0) +
            salary_details.get('Arear Allowance', 0.0) +
            earned_gross
        ])

        total = sum([
            salary_details.get('Provident Fund', 0.0) + salary_details.get('Professional Tax', 0.0) + salary_details.get('Loan', 0.0) + salary_details.get('Labour Welfare Fund', 0.0) +
            salary_details.get('Others', 0.0) + salary_details.get('TDS Deduction', 0.0) + salary_details.get('Employer State Insurance', 0.0)
        ])

        net_pay = (
            total_earnings - total
        )

        fromdate = datetime.strptime(str(s['date_of_birth']), '%Y-%m-%d')
        dob = fromdate.strftime('%d-%m-%Y')
        fromdate = datetime.strptime(str(s['date_of_joining']), '%Y-%m-%d')
        doj = fromdate.strftime('%d-%m-%Y')
        
        
        row=[]
        if emp.employee_category == "ITI":
            index+=1
            row += [
                index, s['employee'], s['employee_name'], dob, emp.grade , doj, s['department'],
                s['designation'],emp.aadhar_no , emp.pan_no, emp.uan_no, emp.esic_no, emp.bank_ac_no, emp.bank_type, emp.employee_category,
                s['total_working_days'], emp.basic or 0, emp.house_rent_allowance or 0, emp.conveyance_allowance or 0,
                emp.education_allowance or 0, emp.food_allowance or 0, emp.attire_allowance or 0, emp.mobile_allowance or 0,
                emp.medical_allowance or 0, emp.special_allowance or 0, emp.fixed_earning, s['total_present_days_of_employee'], holiday_data['total_holiday'],
                coff_count, cl_count, sl_count, el_count, holiday_data['total_other_holiday'], lop_count, s['absent_days'], s['payment_days'],
                round(salary_details['Basic']),
                round(salary_details['House Rent Allowance']), round(salary_details['Conveyance']), round(salary_details['Education Allowance']),
                round(salary_details['Food Allowance']), round(salary_details['Attire Allowance']),round(salary_details['Mobile Allowance']) ,
                round(salary_details['Medical Re imbursement']), round(salary_details['Special Allowance']), round(earned_gross),
                round(salary_details['Attendance Allowance']), round(salary_details['Heat Allowance']),round(salary_details['TL Allowance']),round(salary_details['EL Encashment']), round(salary_details['Gratuity']), round(salary_details['Travel Allowance']),0 , s['ot_hours'],
                round(salary_details['Overtime']),round(salary_details['Arear Allowance']),round(total_earnings), round(salary_details['Provident Fund']) ,
                round(salary_details['Employer State Insurance'],0), round(salary_details['Professional Tax']), round(salary_details['Labour Welfare Fund']),
                round(salary_details['Loan']),round(salary_details['Others']), round(salary_details['TDS Deduction']), round(total,0), round(net_pay)]
            
            data.append(row)

            working_days_total += s['total_working_days']
            fixed_basic_total += emp.basic
            fixed_hra_total += emp.house_rent_allowance
            fixed_conveyance_total +=emp.conveyance_allowance
            fixed_education_total += emp.education_allowance
            fixed_food_total += emp.food_allowance
            fixed_attire_total += emp.attire_allowance or 0
            fixed_mobile_total += emp.mobile_allowance or 0
            fixed_medical_total += emp.medical_allowance
            fixed_special_total += emp.special_allowance
            fixed_gross_total += emp.fixed_earning
            present_days_total += s['total_present_days_of_employee']
            sph_total += holiday_data['total_holiday']
            coff_total += coff_count
            cl_total += cl_count
            sl_total += sl_count
            el_total += el_count
            nf_total += holiday_data['total_other_holiday']
            lop_total += lop_count
            absent_total += s['absent_days']
            npd_total += s['payment_days']
            gross_basic_total += round(salary_details['Basic'])
            gross_hra_total += round(salary_details['House Rent Allowance'])
            gross_conveyance_total += round(salary_details['Conveyance'])
            gross_education_total += round(salary_details['Education Allowance'])
            gross_food_total += round(salary_details['Food Allowance'])
            gross_attire_total += round(salary_details['Attire Allowance'])
            gross_mobile_total += round(salary_details['Mobile Allowance'])
            gross_medical_total += round(salary_details['Medical Re imbursement'])
            gross_special_total += round(salary_details['Special Allowance'])
            earned_gross_total += round(earned_gross)
            attendance_allow_total += round(salary_details['Attendance Allowance'])
            heat_allow_total += round(salary_details['Heat Allowance'])
            tl_allow_total += round(salary_details['TL Allowance'])
            other_allow_total += round(salary_details['Arear Allowance'])
            gratuity_total += round(salary_details['Gratuity'])
            ot_hours += s['ot_hours']
            overtime_total += round(salary_details['Overtime'])
            travel_allow += round(salary_details['Travel Allowance'])
            el_enhancement += round(salary_details['EL Encashment'])
            total_earnings_total += round(total_earnings)
            epf_total += round(salary_details['Provident Fund'])
            esci_total += round(salary_details['Employer State Insurance'],0)
            prof_tax_total += round(salary_details['Professional Tax'])
            loan_total += round(salary_details['Loan'])
            lwf_total += round(salary_details['Labour Welfare Fund'])
            tds_total += round(salary_details['TDS Deduction'])
            others_total += round(salary_details['Others'])
            total_total += round(total)
            net_pay_total += round(net_pay)
            total_arrear_allowance +=round(salary_details['Arear Allowance'])
    row_total += ['Total', '', '', '', '', '', '', '', '', '', '', '', '','','', working_days_total, fixed_basic_total, fixed_hra_total, fixed_conveyance_total,
        fixed_education_total, fixed_food_total,fixed_attire_total,fixed_mobile_total, fixed_medical_total, fixed_special_total, fixed_gross_total, present_days_total, sph_total,
        coff_total,cl_total, sl_total, el_total, nf_total, lop_total, absent_total, npd_total, gross_basic_total, gross_hra_total, gross_conveyance_total,
        gross_education_total, gross_food_total,gross_attire_total, gross_mobile_total, gross_medical_total, gross_special_total, earned_gross_total, attendance_allow_total, heat_allow_total, tl_allow_total, el_enhancement, gratuity_total, travel_allow, 0,
        ot_hours, overtime_total, other_allow_total, total_earnings_total, epf_total, round(esci_total,0), prof_tax_total,  lwf_total, loan_total, others_total, tds_total,  round(total_total), net_pay_total]
    data.append(row_total)
    return data


def check_holiday(date, emp):
    # Get the holiday list for the employee
    holiday_list = frappe.db.get_value('Employee', emp, 'holiday_list')
    
    # Parameterized query to prevent SQL injection and ensure proper handling of input values
    query = """
        SELECT `tabHoliday`.holiday_date, `tabHoliday`.weekly_off
        FROM `tabHoliday List`
        LEFT JOIN `tabHoliday` ON `tabHoliday`.parent = `tabHoliday List`.name
        WHERE `tabHoliday List`.name = %s AND holiday_date = %s
    """
    holiday = frappe.db.sql(query, (holiday_list, date), as_dict=True)
    if holiday:
        if holiday[0].weekly_off == 1:
            return "WW"
        else:
            return "HH"
    return None