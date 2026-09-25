import frappe
import datetime
from frappe.utils import cstr, add_days, date_diff, getdate, format_date
from frappe import _, bold
from frappe.utils.csvutils import UnicodeWriter, read_csv_content
from frappe.utils.file_manager import get_file
from frappe.model.document import Document
from frappe.utils.background_jobs import enqueue
from datetime import date, timedelta, datetime
import openpyxl
from openpyxl import Workbook
from openpyxl.drawing.image import Image as xlImage
import io
import openpyxl
import requests
import xlrd
import re
from openpyxl.styles import Font, Alignment, Border, Side
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.styles import GradientFill, PatternFill
from six import BytesIO, string_types
from openpyxl.styles.numbers import FORMAT_PERCENTAGE
from frappe.utils import (getdate, cint, add_months, date_diff, add_days,format_date,
    nowdate, get_datetime_str, cstr, get_datetime, now_datetime, format_datetime)



@frappe.whitelist()
def download():
    filename = 'Salary register for TT'
    test = build_xlsx_response(filename)

@frappe.whitelist()
def make_xlsx(data, sheet_name=None, wb=None, column_widths=None):
    args = frappe.local.form_dict
    column_widths = column_widths or []
    if wb is None:
        wb = openpyxl.Workbook()
    ws = wb.create_sheet(sheet_name, 0)

    header1 = ['JOHOKU MANUFACTURING PRIVATE LIMITED','','','','','','','','','','','','','','','','','','','','','','','','','','','','Prepared By','','','','Checked By','','','','Verified By','','','','','Aproved By']
    ws.append(header1)
    header2 = [' ']
    ws.append(header2)
    start_date=getdate(args.from_date)
    attmonth=start_date.strftime('%B')
    att_year = start_date.year
    title=['SALARY STATEMENT FOR THE MONTH OF'+ ' '+str(attmonth).upper() + ' '+ str(att_year)+ ' ' +'TEMPROARY TRAINEE']
    ws.append(title)

    
    header3 = ['S.No','Emp Code','Name','DOB','DOJ','Department','Designation','Bank A/C No','Mode','UAN No.','ESIC No.','Category','No.of Working days',
                        'FIXED','','','Present Days','S/PH','COFF','CL','SL','EL','N/F','LOP','Absent','NPD','GROSS EARNINGS','','',
                        'Variable Earnings','','','','','','','','Total Earning','Deductions','','','','','','','','Net Pay INR' ]
    ws.append(header3)
    header4 = ['','','','','','','','','','','','','','Basic','HRA','Fixed Gross','','','','','','','','','','','Basic','HRA','Total Gross Earnings',
                'Attendance allowance','Heat Allowance','Travel Allowance','Bonus','OT Hours','Overtime Amount','Other Allowance','EL Encashment','',
                'EPF Round Off','ESIC','Prof Tax','Canteen','LIC','LWF','Others','Total','']
    ws.append(header4)

    data = get_data(args)
    for row in data:
        ws.append(row)

    for i in range(1,14):
        ws.merge_cells(start_row=4,start_column=i,end_row=5,end_column=i)
           
    for j in range(17,27):
        ws.merge_cells(start_row=4,start_column=j,end_row=5,end_column=j)

    ws.merge_cells(start_row=4,start_column=38,end_row=5,end_column=38)
    ws.merge_cells(start_row=4,start_column=47,end_row=5,end_column=47) 
    ws.merge_cells(start_row=4,start_column=14,end_row=4,end_column=16)
    ws.merge_cells(start_row=4,start_column=27,end_row=4,end_column=29)

    ws.merge_cells(start_row=4,start_column=30,end_row=4,end_column=37)
    ws.merge_cells(start_row=4,start_column=39,end_row=4,end_column=46)

    ws.merge_cells(start_row=1,start_column=1,end_row=1,end_column=28)
    ws.merge_cells(start_row=2,start_column=1,end_row=2,end_column=28)
    ws.merge_cells(start_row=3,start_column=1,end_row=3,end_column=28)

    ws.merge_cells(start_row=1,start_column=29,end_row=1,end_column=32)
    ws.merge_cells(start_row=2,start_column=29,end_row=3,end_column=32)
    ws.merge_cells(start_row=1,start_column=33,end_row=1,end_column=36)
    ws.merge_cells(start_row=2,start_column=33,end_row=3,end_column=36)

    ws.merge_cells(start_row=1,start_column=37,end_row=1,end_column=41)
    ws.merge_cells(start_row=2,start_column=37,end_row=3,end_column=41)
    ws.merge_cells(start_row=1,start_column=42,end_row=1,end_column=47)
    ws.merge_cells(start_row=2,start_column=42,end_row=3,end_column=47)


    image_url="http://157.245.101.198/files/JohokuNew.png"
    image_data = requests.get(image_url).content
    img = xlImage(io.BytesIO(image_data))
    img_cell = ws.cell(row=1, column=6)
    img.width = 75
    img.height = 40
    ws.add_image(img,"F1")
    img.anchor = 'F1' 


    for header in ws.iter_rows(min_row=4, max_row=5,min_col=1, max_col=47):
        for cell in header:
            cell.fill = PatternFill(fgColor='f4f087', fill_type="solid") 
        
    border = Border(left=Side(border_style='thin', color='000000'),
        right=Side(border_style='thin', color='000000'),
        top=Side(border_style='thin', color='000000'),
        bottom=Side(border_style='thin', color='000000'))

    for rows in ws.iter_rows(min_row=1, max_row=ws.max_row, min_col=1,max_col=47):
        for cell in rows:
            cell.border = border

    for total in ws.iter_rows(min_row=6, max_row=ws.max_row,min_col=13, max_col=13):
        for cell in total:
            cell.fill = PatternFill(fgColor="FABF8F", fill_type = "solid")
    
    for present in ws.iter_rows(min_row=6, max_row=ws.max_row,min_col=17, max_col=25):
        for cell in present:
            cell.fill = PatternFill(fgColor="FABF8F", fill_type = "solid")

    for ot in ws.iter_rows(min_row=6, max_row=ws.max_row,min_col=30, max_col=33):
        for cell in ot:
            cell.fill = PatternFill(fgColor="FABF8F", fill_type = "solid")

    for allow in ws.iter_rows(min_row=6, max_row=ws.max_row,min_col=35, max_col=36):
        for cell in allow:
            cell.fill = PatternFill(fgColor="FABF8F", fill_type = "solid")

    for lic in ws.iter_rows(min_row=6, max_row=ws.max_row,min_col=41, max_col=45):
        for cell in lic:
            cell.fill = PatternFill(fgColor="FABF8F", fill_type = "solid")

    ws.merge_cells(start_row=ws.max_row, start_column=1, end_row=ws.max_row, end_column=12)

    for total in ws.iter_rows(min_row=ws.max_row, max_row=ws.max_row,min_col=1, max_col=47):
        for cell in total:
            cell.fill = PatternFill(fgColor="f4f087", fill_type = "solid")

    align_center = Alignment(horizontal='center',vertical='center')

    for align in ws.iter_rows(min_row=1, max_row=ws.max_row,min_col=1, max_col=47):
        for cell in align:
            cell.alignment = align_center

    for cell in ws["1:1"]:
        cell.font = Font(bold=True,size=29)

    for cell in ws["3:3"]:
        cell.font = Font(bold=True,size=18)

    ws.column_dimensions['B'].width = 20 
    ws.column_dimensions['C'].width = 20 
    ws.column_dimensions['D'].width = 20 
    ws.column_dimensions['E'].width = 20 
    ws.column_dimensions['F'].width = 20 
    ws.column_dimensions['G'].width = 20 
    ws.column_dimensions['H'].width = 20 
    ws.column_dimensions['J'].width = 10
    ws.column_dimensions['K'].width = 10
    ws.column_dimensions['Q'].width = 15
    ws.column_dimensions['M'].width = 18
    ws.column_dimensions['P'].width = 15
    ws.column_dimensions['AC'].width = 18
    ws.column_dimensions['AD'].width = 20
    ws.column_dimensions['AE'].width = 18
    ws.column_dimensions['AF'].width = 18
    ws.column_dimensions['AG'].width = 18
    ws.column_dimensions['AH'].width = 15
    ws.column_dimensions['AI'].width = 18
    ws.column_dimensions['AR'].width = 10
    ws.column_dimensions['AS'].width = 18
    ws.column_dimensions['AJ'].width = 18
    ws.column_dimensions['AK'].width = 18
    ws.column_dimensions['AL'].width = 15
    ws.column_dimensions['AM'].width = 15
    ws.column_dimensions['AT'].width = 15
    ws.column_dimensions['AU'].width = 15


    ws.row_dimensions[1].height = 60
    ws.row_dimensions[2].height = 30
    ws.row_dimensions[3].height = 30
    ws.row_dimensions[4].height = 30
    ws.row_dimensions[5].height = 30

         
    xlsx_file=BytesIO()
    wb.save(xlsx_file)
    return xlsx_file

@frappe.whitelist()
def build_xlsx_response(filename):
    xlsx_file = make_xlsx(filename)
    frappe.response['filename'] = filename + '.xlsx'
    frappe.response['filecontent'] = xlsx_file.getvalue()
    frappe.response['type'] = 'binary'

def get_data(args):
    if args.get('from_date') and args.get('to_date'):
        from_date = args.get('from_date')
        to_date = args.get('to_date')
    var = frappe.db.get_all('Salary Slip', {'start_date': from_date, 'end_date': to_date,'salary_structure':'GT-TT Structure'}, ["*"])
    data=[]
    row=[]
    
    total_workingdays=0
    fixed_basic=0
    fixed_house=0
    fixed_total=0
    total_presentday=0
    total_ss=0
    total_tot_coff=0
    total_tot_cll=0
    total_tot_sll=0
    total_tot_earned=0
    total_nh=0
    total_tot_leavepay=0
    total_payment_day=0
    gross_basic=0
    gross_house=0
    gross_total=0
    variable_attendance=0
    variable_heat=0
    variable_travel=0
    variable_bonus=0
    variable_overtime=0
    variable_overtimeamount=0
    variable_otherallow=0
    variable_elencash=0
    total_grosspay=0
    deduction_pfound=0
    deduction_esi=0
    deduction_protax=0
    deduction_canteen=0
    deduction_licform=0
    deduction_labfund=0
    deduction_others=0
    deduction_total_deduction=0
    absent_total = 0
        
    total_netpay=0
    count = 0


    for i in var:   
        emp_cat=frappe.db.get_value("Employee",{'name':i.employee},['employee_category'])
        if emp_cat=='TT':
            emp=frappe.db.get_all('Employee',{'name':i.employee},["*"])
            count+=1
            # presentcount=frappe.db.count('Attendance',{'status':'Present','attendance_date':['between', (from_date, to_date)],'docstatus':['!=', 2],'employee':i.employee})
            # halfdaycount=frappe.db.count('Attendance',{'status':'Half Day','attendance_date':['between', (from_date, to_date)],'docstatus':['!=', 2],'employee':i.employee})
            # hd=0
            # if halfdaycount > 0: 
            #     hd=halfdaycount*0.5
            # else:
            #     hd=0
            # totalpresent=presentcount+hd  




            present_doc_name_list = frappe.db.get_all(
                "Attendance", 
                filters={
                    'employee': i.employee,
                    'status': 'Present',  
                    'attendance_date': ["between", [from_date, to_date]],
                    'docstatus': ["!=", 2], 
                },
                fields=['attendance_date', 'name', 'employee']
            )

            present_days_count = 0
            halfdays_count = 0
            for doc_name in present_doc_name_list:
                attendance_date = doc_name['attendance_date']
                if not check_holiday(attendance_date, i.employee):
                    present_doc_name = frappe.db.get_value(
                        "Attendance", 
                        filters={
                            'employee': i.employee,
                            'status': 'Present',  
                            'attendance_date': attendance_date,
                            'docstatus': ["!=", 2], 
                        },
                        fieldname=['attendance_date', 'name']
                    )
                    if present_doc_name:
                        present_days_count += 1
                    halfday_doc_name = frappe.db.get_value(
                        "Attendance", 
                        filters={
                            'employee': i.employee,
                            'status': 'Half Day',  
                            'attendance_date': attendance_date,
                            'docstatus': ["!=", 2], 
                        },
                        fieldname=['attendance_date', 'name']
                    )
                    if halfday_doc_name:
                        halfdays_count += 1

            totalpresent = present_days_count + (halfdays_count / 2)

            cof=frappe.db.get_all('Leave Application' , {'from_date':['between',(from_date, to_date)],'to_date':['between',(from_date,to_date)],'leave_type':'Compensatory Off','employee':i.employee},["*"])
            tot_coff=0
            for co in cof:
                tot_coff+=co.total_leave_days
            cll=frappe.db.get_all('Leave Application' , {'from_date':['between',(from_date, to_date)],'to_date':['between',(from_date,to_date)],'leave_type':'Casual Leave','employee':i.employee},["*"])
            tot_cll=0
            for cl in cll:
                tot_cll+=cl.total_leave_days
            sll=frappe.db.get_all('Leave Application' , {'from_date':['between',(from_date, to_date)],'to_date':['between',(from_date,to_date)],'leave_type':'Sick Leave','employee':i.employee},["*"])
            tot_sll=0
            for sl in sll:
                tot_sll+=sl.total_leave_days
            earned=frappe.db.get_all('Leave Application' , {'from_date':['between',(from_date, to_date)],'to_date':['between',(from_date,to_date)],'leave_type':'Earned Leave','employee':i.employee},["*"])
            tot_earned=0
            for ea in earned:
                tot_earned+=ea.total_leave_days
            leavepay=frappe.db.get_all('Leave Application' , {'from_date':['between',(from_date, to_date)],'to_date':['between',(from_date,to_date)],'leave_type':'Leave Without Pay','employee':i.employee},["*"])
            tot_leavepay=0
            for lw in leavepay:
                tot_leavepay+=lw.total_leave_days
            for j in emp:
                basic=frappe.get_value("Salary Detail",{'parent':i.name,'salary_component':'Basic'},['amount']) or 0.0
                house=frappe.get_value("Salary Detail",{'parent':i.name,'salary_component':'House Rent Allowance'},['amount']) or 0.0
                heat=frappe.get_value("Salary Detail",{'parent':i.name,'salary_component':'Heat Allowance'},['amount']) or 0.0
                travel=frappe.get_value("Salary Detail",{'parent':i.name,'salary_component':'Travel Allowance'},['amount']) or 0.0
                bonus=frappe.get_value("Salary Detail",{'parent':i.name,'salary_component':'Bonus'},['amount']) or 0.0
                otherallow=frappe.get_value("Salary Detail",{'parent':i.name,'salary_component':'Other Allowances'},['amount']) or 0.0
                overtime=frappe.get_value("Salary Detail",{'parent':i.name,'salary_component':'Overtime'},['amount']) or 0.0
                attendance=frappe.get_value("Salary Detail",{'parent':i.name,'salary_component':'Attendance Allowance'},['amount']) or 0.0
                pfound=frappe.get_value("Salary Detail",{'parent':i.name,'salary_component':'Provident Fund'},['amount']) or 0.0
                esi=frappe.get_value("Salary Detail",{'parent':i.name,'salary_component':'Employer State Insurance'},['amount']) or 0.0
                protax=frappe.get_value("Salary Detail",{'parent':i.name,'salary_component':'Professional Tax'},['amount']) or 0.0
                labfund=frappe.get_value("Salary Detail",{'parent':i.name,'salary_component':'Labour Welfare Fund'},['amount']) or 0.0
                others=frappe.get_value("Salary Detail",{'parent':i.name,'salary_component':'Others'},['amount']) or 0.0
                ss=check_holiday_ss(from_date,to_date,i.employee)
                nh=check_holiday_nh(from_date,to_date,i.employee)
                ss_days=int(ss)
                nh_days=int(nh)
                tot_fixed=j.basic+j.house_rent_allowance
                tot_gross=basic+house
                overtimeamount=0
                elencash=0
                canteen=0
                licform=0
            

                row=([count, 
                j.employee,
                i.employee_name,
                j.date_of_birth,
                j.date_of_joining,  
                i.department, 
                i.designation,
                i.bank_account_no, 
                i.salary_mode or "", 
                i.una_no, '',
                j.employee_category,
                i.total_working_days,
                j.basic,
                j.house_rent_allowance,
                tot_fixed,
                totalpresent,ss_days,tot_coff,tot_cll,tot_sll,tot_earned,nh_days,tot_leavepay,i['absent_days'],i.payment_days,
                basic,house,
                tot_gross,
                attendance,heat,travel,bonus,overtime,overtimeamount,otherallow,elencash,
                i.gross_pay, pfound,esi, protax,canteen,licform,labfund,others,
                i.total_deduction,
                i.net_pay
                ])    
                total_workingdays+=int(i.total_working_days)	
                fixed_basic+=float(j.basic)
                fixed_house+=float(j.house_rent_allowance)
                fixed_total+=tot_fixed
                total_presentday+=float(totalpresent)
                total_ss+=ss_days
                total_tot_coff+=tot_coff
                total_tot_cll+=tot_cll
                total_tot_sll+=tot_sll
                total_tot_earned+=tot_earned
                total_nh+=nh_days
                total_tot_leavepay+=tot_leavepay
                absent_total += i['absent_days']
                total_payment_day+=i.payment_days
                gross_basic+=int(basic)
                gross_house+=int(house)
                gross_total+=tot_gross
                variable_attendance+=attendance
                variable_heat+=heat
                variable_travel+=travel
                variable_bonus+=bonus
                variable_overtime+=overtime
                variable_overtimeamount+=overtimeamount
                variable_otherallow+=otherallow
                variable_elencash+=int(elencash)
                total_grosspay+=float(i.gross_pay)
                deduction_pfound+=int(pfound)
                deduction_esi+=float(esi)
                deduction_protax+=int(protax)
                deduction_canteen+=int(canteen)
                deduction_licform+=int(licform)
                deduction_labfund+=int(labfund)
                deduction_others+=int(others)
                deduction_total_deduction+=i.total_deduction
                total_netpay+=float(i.net_pay)

            data.append(row)

    row1=[ 'TOTAL','','','','','','','','','','','',total_workingdays,fixed_basic,fixed_house,fixed_total,total_presentday,
        total_ss, total_tot_coff, total_tot_cll, total_tot_sll,total_tot_earned,total_nh, total_tot_leavepay,absent_total,
        total_payment_day, gross_basic,gross_house,gross_total,variable_attendance, variable_heat,variable_travel,variable_bonus,
        variable_overtime, variable_overtimeamount, variable_otherallow, variable_elencash, total_grosspay,deduction_pfound,
        deduction_esi, deduction_protax, deduction_canteen, deduction_licform, deduction_labfund,deduction_others,
        deduction_total_deduction, total_netpay]
    data.append(row1)
    return data                

@frappe.whitelist()
def get_dates(from_date,to_date):
    no_of_days = date_diff(add_days(to_date, 1), from_date)
    dates = [add_days(from_date, i) for i in range(0, no_of_days)]
    return dates

@frappe.whitelist()
def check_holiday_ss(from_date, to_date, emp):
    dates = get_dates(from_date, to_date)
    tot = 0
    for date in dates:
        holiday_list = frappe.db.get_value('Employee', {'name': emp}, 'holiday_list')
        holiday = frappe.db.sql("""
            SELECT tabHoliday.holiday_date, tabHoliday.weekly_off
            FROM `tabHoliday List` AS holiday_list
            LEFT JOIN `tabHoliday` ON tabHoliday.parent = holiday_list.name
            WHERE holiday_list.name = %s AND holiday_date = %s
        """, (holiday_list, date), as_dict=True)
        
        doj = frappe.db.get_value("Employee", {'name': emp}, "date_of_joining")
        
        if holiday:
            if doj < holiday[0].holiday_date:
                if holiday[0].weekly_off == 1 or holiday[0].holiday_type == 'PH':
                    tot += 1
    return tot

@frappe.whitelist()
def check_holiday_nh(from_date, to_date, emp):
    dates = get_dates(from_date, to_date)
    tot = 0
    for date in dates:
        holiday_list = frappe.db.get_value('Employee', {'name': emp}, 'holiday_list')
        holiday = frappe.db.sql("""
            SELECT tabHoliday.holiday_date, tabHoliday.weekly_off
            FROM `tabHoliday List` AS holiday_list
            LEFT JOIN `tabHoliday` ON tabHoliday.parent = holiday_list.name
            WHERE holiday_list.name = %s AND holiday_date = %s
        """, (holiday_list, date), as_dict=True)
        
        doj = frappe.db.get_value("Employee", {'name': emp}, "date_of_joining")
        
        if holiday:
            if doj < holiday[0].holiday_date:
                if holiday[0].holiday_type == 'NH' or holiday[0].holiday_type == 'FH':
                    tot += 1
    return tot



def check_holiday(date, emp):
    # Get the holiday list for the employee
    holiday_list = frappe.db.get_value('Employee', emp, 'holiday_list')
    
    # Parameterized query to prevent SQL injection and ensure proper handling of input values
    query = """
        SELECT `tabHoliday`.holiday_date, `tabHoliday`.weekly_off
        FROM `tabHoliday List`
        LEFT JOIN `tabHoliday` ON `tabHoliday`.parent = `tabHoliday List`.name
        WHERE `tabHoliday List`.name = %s AND holiday_date = %s
    """
    holiday = frappe.db.sql(query, (holiday_list, date), as_dict=True)
    if holiday:
        if holiday[0].weekly_off == 1:
            return "WW"
        else:
            return "HH"
    return None