# Copyright (c) 2024, TEAMPRO and contributors
# For license information, please see license.txt

# import frappe
from frappe.model.document import Document
import frappe
import openpyxl
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side
from openpyxl.styles import GradientFill, PatternFill
from six import BytesIO
import io  # Add this import to resolve the NameError

from datetime import datetime
from datetime import datetime, timedelta
from frappe.utils import (getdate, cint, add_months, date_diff, add_days, format_date,
    nowdate, get_datetime_str, cstr, get_datetime, now_datetime, format_datetime)
from openpyxl.drawing.image import Image as xlImage
import requests


# class SalaryRegisterReports(Document):
# 	pass


@frappe.whitelist()
def download():
    filename = 'Salary Summary.xlsx'
    build_xlsx_response(filename)

def make_xlsx(sheet_name=None):
    args = frappe.local.form_dict
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_name if sheet_name else 'Sheet1'
    
    data = get_data(args)
    
    if not data:  # Check if data is empty or None
        frappe.throw("No data available to generate the report.")
    
    
    image_url="http://157.245.101.198/files/JohokuNew.png"
    image_data = requests.get(image_url).content
    img = xlImage(io.BytesIO(image_data))
    img_cell = ws.cell(row=1, column=6)
    img.width = 100
    img.height = 80
    ws.add_image(img,"E1")
         
        
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=6)
    ws["G1"].value = "JOHOKU MANUFACTURING PRIVATE LIMITED"
    ws["G1"].alignment = Alignment(horizontal='left', vertical='center', wrap_text=True)
    ws["G1"].font = Font(bold=True, size=14)
    ws.merge_cells(start_row=1, start_column=7, end_row=1, end_column=18)
    for row in data:
        if not isinstance(row, (list, tuple)):
            row = [row]  
        ws.append(row)

    
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=18)
    for x in range(1,12):
        ws.merge_cells(start_row=3, start_column=x, end_row=5, end_column=x)
    ws.merge_cells(start_row=3, start_column=12, end_row=3, end_column=16)
    for x in range(12,17):
        ws.merge_cells(start_row=4, start_column=x, end_row=5, end_column=x)
    ws.merge_cells(start_row=3, start_column=17, end_row=5, end_column=17)
    ws.merge_cells(start_row=3, start_column=18, end_row=5, end_column=18)
    apply_styles(ws)
    set_column_widths(ws)

    xlsx_file = BytesIO()
    wb.save(xlsx_file)
    xlsx_file.seek(0)
    return xlsx_file

def apply_styles(ws):
    align_center = Alignment(horizontal='center', vertical='center', wrap_text=True)
    align_right = Alignment(horizontal='right', vertical='top', wrap_text=True)
    align_left = Alignment(horizontal='left', vertical='top', wrap_text=True)
    header_font = Font(bold=True, size=14)
    text_font = Font(bold=True, size=8)
    text_font_header = Font(bold=True, size=10)
    text_font_data = Font(bold=False, size=9)
    border = Border(
        left=Side(border_style='thin'),
        right=Side(border_style='thin'),
        top=Side(border_style='thin'),
        bottom=Side(border_style='thin')
    )
    outline_border = Border(
    left=thin_border,
    right=thin_border,
    top=thin_border,
    bottom=thin_border
)

# Apply the outline border for the entire range
    outline_range = ws.iter_rows(min_row=1, max_row=ws.max_row, min_col=1, max_col=18)
    for rows in outline_range:
        for cell in rows:
            cell.border = outline_border
                cell.border = border
    for rows in ws.iter_rows(min_row=2, max_row=2, min_col=1, max_col=18):
        for cell in rows:
            cell.font = header_font
            cell.alignment = align_center
            cell.border = border
    for x in range(3,6):
        for rows in ws.iter_rows(min_row=x, max_row=x, min_col=1, max_col=18):
            for cell in rows:
                cell.font = text_font_header
                cell.alignment = align_center
                cell.border = border
    for x in range(6,ws.max_row + 1):
        for rows in ws.iter_rows(min_row=x, max_row=x, min_col=1, max_col=18):
            for cell in rows:
                cell.font = text_font_data
                cell.alignment = align_right
                cell.border = border
    
    

def set_column_widths(ws):
    column_widths = [10] *18 
    for i, width in enumerate(column_widths, start=1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = width
    row_heights = [60, 60, 20,20,20] + [20] * (ws.max_row - 1) +[30]*1 
    for i, height in enumerate(row_heights, start=1):
        ws.row_dimensions[i].height = height

def build_xlsx_response(filename):
    xlsx_file = make_xlsx(sheet_name=filename)
    frappe.response['filename'] = filename + '.xlsx'
    frappe.response['filecontent'] = xlsx_file.getvalue()
    frappe.response['type'] = 'binary'

def get_salary_component_data(slip_name_list):
    basic_amount = ot_amount = attendance_amount = heat_allowance_amount = travel_llowance_amount = el_amount = other_allowances_amount = pf_amount = 0
    esic_amount = pt_amount = loan_amount = others_amount = 0
    components = ['Basic', 'Overtime', 'Attendance Allowance', 'Heat Allowance','Travel Allowance','EL Encashment','Other Allowances',
        'Provident Fund','Employer State Insurance', 'Professional Tax', 'Loan', 'Others']
    for slip_name in slip_name_list:
        for comp in components:
            if comp == "Basic":
                basic = frappe.db.sql(""" SELECT SUM(amount) AS salary_amt FROM `tabSalary Detail` WHERE parent = %s AND salary_component = %s """, (slip_name, comp), as_dict=True)
                basic_amount += basic[0]['salary_amt'] or 0
            elif comp == "Overtime":
                ot = frappe.db.sql(""" SELECT SUM(amount) AS salary_amt FROM `tabSalary Detail` WHERE parent = %s AND salary_component = %s """, (slip_name, comp), as_dict=True)
                ot_amount += ot[0]['salary_amt'] or 0
            elif comp == "Attendance Allowance":
                attendance = frappe.db.sql(""" SELECT SUM(amount) AS salary_amt FROM `tabSalary Detail` WHERE parent = %s AND salary_component = %s """, (slip_name, comp), as_dict=True)
                attendance_amount += attendance[0]['salary_amt'] or 0
            elif comp == "Heat Allowance":
                heat_allowance = frappe.db.sql(""" SELECT SUM(amount) AS salary_amt FROM `tabSalary Detail` WHERE parent = %s AND salary_component = %s """, (slip_name, comp), as_dict=True)
                heat_allowance_amount += heat_allowance[0]['salary_amt'] or 0
            elif comp == "Travel Allowance":
                travel_llowance = frappe.db.sql(""" SELECT SUM(amount) AS salary_amt FROM `tabSalary Detail` WHERE parent = %s AND salary_component = %s """, (slip_name, comp), as_dict=True)
                travel_llowance_amount += travel_llowance[0]['salary_amt'] or 0
            elif comp == "EL Encashment":
                el = frappe.db.sql(""" SELECT SUM(amount) AS salary_amt FROM `tabSalary Detail` WHERE parent = %s AND salary_component = %s """, (slip_name, comp), as_dict=True)
                el_amount += el[0]['salary_amt'] or 0
            elif comp == "Other Allowances":
                other_allowances = frappe.db.sql(""" SELECT SUM(amount) AS salary_amt FROM `tabSalary Detail` WHERE parent = %s AND salary_component = %s """, (slip_name, comp), as_dict=True)
                other_allowances_amount += other_allowances[0]['salary_amt'] or 0
            elif comp == "Provident Fund":
                pf = frappe.db.sql(""" SELECT SUM(amount) AS salary_amt FROM `tabSalary Detail` WHERE parent = %s AND salary_component = %s """, (slip_name, comp), as_dict=True)
                pf_amount += pf[0]['salary_amt'] or 0
            elif comp == "Employer State Insurance":
                esic = frappe.db.sql(""" SELECT SUM(amount) AS salary_amt FROM `tabSalary Detail` WHERE parent = %s AND salary_component = %s """, (slip_name, comp), as_dict=True)
                esic_amount += esic[0]['salary_amt'] or 0
            elif comp == "Professional Tax":
                pt = frappe.db.sql(""" SELECT SUM(amount) AS salary_amt FROM `tabSalary Detail` WHERE parent = %s AND salary_component = %s """, (slip_name, comp), as_dict=True)
                pt_amount += pt[0]['salary_amt'] or 0
            elif comp == "Loan":
                loan = frappe.db.sql(""" SELECT SUM(amount) AS salary_amt FROM `tabSalary Detail` WHERE parent = %s AND salary_component = %s """, (slip_name, comp), as_dict=True)
                loan_amount += loan[0]['salary_amt'] or 0
            elif comp == "Others":
                others = frappe.db.sql(""" SELECT SUM(amount) AS salary_amt FROM `tabSalary Detail` WHERE parent = %s AND salary_component = %s """, (slip_name, comp), as_dict=True)
                others_amount += others[0]['salary_amt'] or 0
    data = [basic_amount,ot_amount,attendance_amount,heat_allowance_amount,travel_llowance_amount,el_amount,other_allowances_amount,pf_amount,
    esic_amount,pt_amount,loan_amount,others_amount]
    return data

def get_data(args):
# def get_data():
#     from_date ="2024-10-21"
#     to_date ="2024-11-20"
    from_date = args.get('from_date')
    to_date = args.get('to_date')

    if not from_date or not to_date:
        frappe.throw("From Date and To Date are required.")
    from_date_str =from_date

    if isinstance(from_date_str, str):
        month_year_format_str = datetime.strptime(from_date_str, '%Y-%m-%d')
    else:
        month_year_format_str = from_date_str  

    month_year_format = month_year_format_str.strftime('%B %Y')
    data = [
            ["SALARY STATEMENT FOR THE MONTH OF " + month_year_format],
            ["Category","No of Employees","Basic","Gross Salary","OT","Attendance Allowance","Heat Allowance","TL Allowance",
            "EL Encashment","Other Allowance","Total Earning","Deductions","","","","","Total Deduction","Net Salary"],
            [""]*11 + ["EPF - Employee Contribution","ESIC","PT","Loan","Others"],
            [""],
            [""]
        ]
    categories = ['Staff', 'DT', 'ITI', 'GT', 'TT']
    total_basic = total_employee_count = 0
    total_gross_salary = total_ot = total_att = total_heat = total_tl = total_el = total_other_allow = total_pf = total_esi = total_pt = total_loan =0
    total_others = total_deduction_salary = total_gross_salary = total_net_pay_salary = total_sum_of_gross_pay =0
    row =[]
    for category in categories:
        if category == 'Staff':
            employee_category = "Staff"
        elif category == 'DT':
            employee_category = "DT Trainee"
        elif category == 'ITI':
            employee_category = "ITI Trainee"
        elif category == 'GT':
            employee_category = "GT Trainee"
        elif category == 'TT':
            employee_category = "TT Trainee"
        slip_name_list = []
        salary_slips = frappe.db.sql(""" SELECT * FROM `tabSalary Slip` WHERE start_date = %s AND end_date = %s AND docstatus != 2 AND employee_category = %s""", (from_date, to_date, category), as_dict=True)
        salary_slip_data = frappe.db.sql(""" SELECT COUNT(employee) AS employee_count, SUM(gross_pay) AS sum_of_gross_pay,SUM(total_deduction) AS total_deduction,SUM(net_pay) AS total_net_pay FROM `tabSalary Slip` WHERE start_date = %s AND end_date = %s AND docstatus != 2 AND employee_category = %s""", (from_date, to_date,category), as_dict=True)

        employee_count = salary_slip_data[0].employee_count
        sum_of_gross_pay = salary_slip_data[0].sum_of_gross_pay
        total_deduction = salary_slip_data[0].total_deduction
        total_net_pay = salary_slip_data[0].total_net_pay
        total_gross = 0
        for s in salary_slips:
            if not isinstance(s, dict):
                continue  
            slip_name_list.append(s['name'])
            employee_data = frappe.db.sql(""" SELECT fixed_earning FROM `tabEmployee` WHERE name = %s""", (s['employee']), as_dict=True)
            fixed_earning = employee_data[0].fixed_earning
            total_gross += fixed_earning
            # print(employee_data)
        staff_data = get_salary_component_data(slip_name_list)
        # print(total_gross)
        # print(sum_of_gross_pay)    
        # print(employee_category)
        total_employee_count += employee_count
        total_basic += staff_data[0]
        total_gross_salary += total_gross
        total_ot += staff_data[1]
        total_att += staff_data[2]
        total_heat += staff_data[3]
        total_tl += staff_data[4]
        total_el += staff_data[5]
        total_other_allow += staff_data[6]
        total_pf += staff_data[7]
        total_esi += staff_data[8]
        total_pt += staff_data[9]
        total_loan += staff_data[10]
        total_others += staff_data[11]
        total_deduction_salary += total_deduction
        total_net_pay_salary += total_net_pay
        total_sum_of_gross_pay += sum_of_gross_pay
        row =[
            employee_category,  
            employee_count,
            staff_data[0],
            total_gross,
            staff_data[1],
            staff_data[2], 
            staff_data[3],
            staff_data[4],
            staff_data[5],
            staff_data[6],
            sum_of_gross_pay, 
            staff_data[7],
            staff_data[8],
            staff_data[9],
            staff_data[10],
            staff_data[11],
            total_deduction,
            total_net_pay]
        data.append(row)
        # print (row)
    # print(total_basic)
    list_of_total = ["Total", total_employee_count,total_basic,
    total_gross_salary, total_ot, total_att, total_heat, total_tl, total_el, total_other_allow,total_sum_of_gross_pay, total_pf,
    total_esi, total_pt, total_loan,total_others, total_deduction_salary, total_net_pay_salary
]     
    data.append(list_of_total)
    return data

