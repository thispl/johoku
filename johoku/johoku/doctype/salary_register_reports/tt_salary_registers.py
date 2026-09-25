import frappe
import datetime
from frappe.utils import cstr, add_days, date_diff, getdate, format_date
from frappe import _, bold
from frappe.utils.csvutils import UnicodeWriter, read_csv_content
from frappe.utils.file_manager import get_file
from frappe.model.document import Document
from frappe.utils.background_jobs import enqueue
from datetime import date, timedelta, datetime
import openpyxl
from openpyxl import Workbook
from openpyxl.drawing.image import Image as xlImage
import io
import openpyxl
import requests
import xlrd
import re
from openpyxl.styles import Font, Alignment, Border, Side
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.styles import GradientFill, PatternFill
from six import BytesIO, string_types
from openpyxl.styles.numbers import FORMAT_PERCENTAGE
from frappe.utils import (getdate, cint, add_months, date_diff, add_days,format_date,
	nowdate, get_datetime_str, cstr, get_datetime, now_datetime, format_datetime)



@frappe.whitelist()
def download():
	filename = 'Salary register for TT'
	test = build_xlsx_response(filename)

@frappe.whitelist()
def make_xlsx(data, sheet_name=None, wb=None, column_widths=None):
	args = frappe.local.form_dict
	column_widths = column_widths or []
	if wb is None:
		wb = openpyxl.Workbook()
	ws = wb.create_sheet(sheet_name, 0)

	header1 = ['JOHOKU MANUFACTURING PRIVATE LIMITED','','','','','','','','','','','','','','','','','','','','','','','','','','','','','','','','','','Prepared By','','','Checked By','','','','','','','Verified By','','','','','Approved By']
	ws.append(header1)
	header2 = [' ']
	ws.append(header2)
	end_date=getdate(args.to_date)
	attmonth=end_date.strftime('%B')
	att_year = end_date.year
	title=['SALARY STATEMENT FOR THE MONTH OF'+ ' '+str(attmonth).upper() + ' '+ str(att_year)+ ' ' +'TEMPROARY TRAINEE']
	ws.append(title)

	
	header3 = ['S.No','Emp Code','Name','DOB','Grade','DOJ','Department','Designation','Aadhar No','PAN No','UAN No.','ESIC No.','Bank A/C No','Bank Type','Category','No Working day',
						'Fixed Gross','','','','Present Days','S/PH','COFF','CL','SL','EL','N/F','LOP','Absent','NPD','GROSS EARNINGS','','','',
						'Variable Earnings','','','','','','','','','','Total Earnings','Deductions','','','','','','','Total Deductions','Net Pay INR' ]
	ws.append(header3)
	header4 = ['','','','','','','','','','','','','','','','','Basic','HRA','Spl Allow','Total Fixed Gross Earning','','','','','','','','','','','Basic','HRA','Spl Allow','Total Gross Earnings',
				'Attendance Bonus','Heat Allowance','TL Allowance','EL Encashment','Gratuity Pay','Travel Allowance','Bonus','OT Hrs','OT Amount','Arear Allowance','',
				'EPF Round Off','ESIC Round up off','PT','LWF','Loan','Other Deduction','TDS Deduction','','']
	ws.append(header4)

	data = get_data(args)
	for row in data:
		ws.append(row)

	for i in range(1,17):
		ws.merge_cells(start_row=4,start_column=i,end_row=5,end_column=i)
		   
	for j in range(21,31):
		ws.merge_cells(start_row=4,start_column=j,end_row=5,end_column=j)

	ws.merge_cells(start_row=4,start_column=46,end_row=4,end_column=52)
	ws.merge_cells(start_row=4,start_column=45,end_row=5,end_column=45) 
	ws.merge_cells(start_row=4,start_column=17,end_row=4,end_column=20)
	ws.merge_cells(start_row=4,start_column=31,end_row=4,end_column=34)

	ws.merge_cells(start_row=4,start_column=35,end_row=4,end_column=44)
	ws.merge_cells(start_row=4,start_column=53,end_row=5,end_column=53)
	ws.merge_cells(start_row=4,start_column=54,end_row=5,end_column=54)

	ws.merge_cells(start_row=1,start_column=1,end_row=1,end_column=34)
	ws.merge_cells(start_row=2,start_column=1,end_row=2,end_column=34)
	ws.merge_cells(start_row=3,start_column=1,end_row=3,end_column=34)

	ws.merge_cells(start_row=1,start_column=35,end_row=1,end_column=37)
	ws.merge_cells(start_row=2,start_column=35,end_row=3,end_column=37)
	ws.merge_cells(start_row=1,start_column=38,end_row=1,end_column=44)
	ws.merge_cells(start_row=2,start_column=38,end_row=3,end_column=40)
	ws.merge_cells(start_row=2,start_column=41,end_row=3,end_column=44)
	ws.merge_cells(start_row=1,start_column=45,end_row=1,end_column=49)
	ws.merge_cells(start_row=2,start_column=45,end_row=3,end_column=49)
	ws.merge_cells(start_row=1,start_column=50,end_row=1,end_column=54)
	ws.merge_cells(start_row=2,start_column=50,end_row=3,end_column=54)


	image_url="http://157.245.101.198/files/JohokuNew.png"
	# image_data = requests.get(image_url).content
	image_data = requests.get(image_url, verify=False).content
	img = xlImage(io.BytesIO(image_data))
	img_cell = ws.cell(row=1, column=6)
	img.width = 75
	img.height = 40
	ws.add_image(img,"F1")
	img.anchor = 'F1' 


	for header in ws.iter_rows(min_row=4, max_row=5,min_col=1, max_col=54):
		for cell in header:
			cell.fill = PatternFill(fgColor='f4f087', fill_type="solid") 
			cell.font = Font(bold=True,size=10)
		
	border = Border(left=Side(border_style='thin', color='000000'),
		right=Side(border_style='thin', color='000000'),
		top=Side(border_style='thin', color='000000'),
		bottom=Side(border_style='thin', color='000000'))

	for rows in ws.iter_rows(min_row=1, max_row=ws.max_row, min_col=1,max_col=54):
		for cell in rows:
			cell.border = border

	for total in ws.iter_rows(min_row=6, max_row=ws.max_row,min_col=16, max_col=16):
		for cell in total:
			cell.fill = PatternFill(fgColor="FABF8F", fill_type = "solid")
	
	for present in ws.iter_rows(min_row=6, max_row=ws.max_row,min_col=21, max_col=29):
		for cell in present:
			cell.fill = PatternFill(fgColor="FABF8F", fill_type = "solid")

	for ot in ws.iter_rows(min_row=6, max_row=ws.max_row,min_col=35, max_col=41):
		for cell in ot:
			cell.fill = PatternFill(fgColor="FABF8F", fill_type = "solid")

	for allow in ws.iter_rows(min_row=6, max_row=ws.max_row,min_col=43, max_col=43):
		for cell in allow:
			cell.fill = PatternFill(fgColor="FABF8F", fill_type = "solid")

	for lic in ws.iter_rows(min_row=6, max_row=ws.max_row,min_col=52, max_col=52):
		for cell in lic:
			cell.fill = PatternFill(fgColor="FABF8F", fill_type = "solid")

	ws.merge_cells(start_row=ws.max_row, start_column=1, end_row=ws.max_row, end_column=15)

	for total in ws.iter_rows(min_row=ws.max_row, max_row=ws.max_row,min_col=1, max_col=54):
		for cell in total:
			cell.fill = PatternFill(fgColor="f4f087", fill_type = "solid")

	align_center = Alignment(horizontal='center',vertical='center')

	for align in ws.iter_rows(min_row=1, max_row=ws.max_row,min_col=1, max_col=54):
		for cell in align:
			cell.alignment = align_center

	for cell in ws["1:1"]:
		cell.font = Font(bold=True,size=29)

	for cell in ws["3:3"]:
		cell.font = Font(bold=True,size=18)

	ws.column_dimensions['B'].width = 20 
	ws.column_dimensions['C'].width = 20 
	ws.column_dimensions['D'].width = 20 
	ws.column_dimensions['E'].width = 20 
	ws.column_dimensions['F'].width = 20 
	ws.column_dimensions['G'].width = 20 
	ws.column_dimensions['H'].width = 20 
	ws.column_dimensions['J'].width = 13
	ws.column_dimensions['K'].width = 13
	ws.column_dimensions['Q'].width = 15
	ws.column_dimensions['M'].width = 18
	ws.column_dimensions['P'].width = 15
	ws.column_dimensions['I'].width = 17
	ws.column_dimensions['L'].width = 13
	ws.column_dimensions['T'].width = 15
	ws.column_dimensions['U'].width = 13
	ws.column_dimensions['AC'].width = 10
	ws.column_dimensions['AD'].width = 10
	ws.column_dimensions['AE'].width = 10
	ws.column_dimensions['AF'].width = 10
	ws.column_dimensions['AG'].width = 10
	ws.column_dimensions['AH'].width = 17
	ws.column_dimensions['AQ'].width = 15
	ws.column_dimensions['AI'].width = 18
	ws.column_dimensions['AR'].width = 13
	ws.column_dimensions['AS'].width = 18
	ws.column_dimensions['AJ'].width = 18
	ws.column_dimensions['AK'].width = 18
	ws.column_dimensions['AL'].width = 15
	ws.column_dimensions['AM'].width = 15
	ws.column_dimensions['AN'].width = 15
	ws.column_dimensions['AT'].width = 15
	ws.column_dimensions['AU'].width = 15
	ws.column_dimensions['BB'].width = 10
	ws.column_dimensions['BA'].width = 15
	ws.column_dimensions['AZ'].width = 15


	ws.row_dimensions[1].height = 60
	ws.row_dimensions[2].height = 30
	ws.row_dimensions[3].height = 30
	ws.row_dimensions[4].height = 30
	ws.row_dimensions[5].height = 30

		 
	xlsx_file=BytesIO()
	wb.save(xlsx_file)
	return xlsx_file

@frappe.whitelist()
def build_xlsx_response(filename):
	xlsx_file = make_xlsx(filename)
	frappe.response['filename'] = filename + '.xlsx'
	frappe.response['filecontent'] = xlsx_file.getvalue()
	frappe.response['type'] = 'binary'


def get_leave_count(employee, leave_type, from_date, to_date):
	if not leave_type:
		frappe.throw("Leave Type is required.")

	# Ensure date format consistency (use getdate if necessary)
	from_date = getdate(from_date)
	to_date = getdate(to_date)
	full_day_leave_count = halfdays_leave_count = leave_count = 0
	full_day_leave_count = frappe.db.count("Attendance", {
		'employee': employee,
		'leave_type': leave_type,
		'status': 'On Leave',
		'leave_application':['!=',''],
		'attendance_date': ['between', [from_date, to_date]],
		'docstatus':["!=",2],
	})
	halfdays_leave_count = frappe.db.count("Attendance", {
		'employee': employee,
		'leave_type': leave_type,
		'status': 'Half Day',
		'leave_application':['!=',''],
		'attendance_date': ['between', [from_date, to_date]],
		'docstatus':["!=",2],
	})
	leave_count = full_day_leave_count + (halfdays_leave_count / 2)
	return leave_count

def get_data(args):
	if args.get('from_date') and args.get('to_date'):
		from_date = args.get('from_date')
		to_date = args.get('to_date')
	var = frappe.db.get_all('Salary Slip', {'start_date': from_date, 'end_date': to_date,'salary_structure':'GT-TT Structure'}, ["*"],order_by= 'employee')
	data=[]
	row=[]
	
	total_workingdays=0
	fixed_basic=0
	fixed_house=0
	fixed_special=0
	fixed_total=0
	total_presentday=0
	total_ss=0
	total_tot_coff=0
	total_tot_cll=0
	total_tot_sll=0
	total_tot_earned=0
	total_nh=0
	total_tot_leavepay=0
	total_payment_day=0
	gross_basic=0
	gross_house=0
	gross_special=0
	gross_total=0
	variable_attendance=0
	variable_heat=0
	variable_tl=0
	variable_gratuity=0
	variable_other=0
	variable_travel=0
	variable_bonus=0
	variable_overtime=0
	variable_overtimeamount=0
	variable_arrearallow=0
	variable_elencash=0
	total_grosspay=0
	deduction_pfound=0
	deduction_esi=0
	deduction_protax=0
	deduction_canteen=0
	deduction_licform=0
	deduction_labfund=0
	deduction_loan=0
	deduction_others=0
	deduction_tds=0
	deduction_total_deduction=0
	absent_total = 0
		
	total_netpay=0
	count = 0


	for i in var:   
		emp_cat=frappe.db.get_value("Employee",{'name':i.employee},['employee_category'])
		if emp_cat=='TT':
			emp=frappe.db.get_all('Employee',{'name':i.employee},["*"],order_by= 'name')
			count+=1
			# presentcount=frappe.db.count('Attendance',{'status':'Present','attendance_date':['between', (from_date, to_date)],'docstatus':['!=', 2],'employee':i.employee})
			# halfdaycount=frappe.db.count('Attendance',{'status':'Half Day','attendance_date':['between', (from_date, to_date)],'docstatus':['!=', 2],'employee':i.employee})
			# hd=0
			# if halfdaycount > 0: 
			#     hd=halfdaycount*0.5
			# else:
			#     hd=0
			# totalpresent=presentcount+hd  




			present_doc_name_list = frappe.db.get_all(
				"Attendance", 
				filters={
					'employee': i.employee,
					'status': ['in',['Present','Half Day']],  
					'attendance_date': ["between", [from_date, to_date]],
					'docstatus': ["!=", 2], 
				},
				fields=['attendance_date', 'name', 'employee']
			)

			present_days_count_ = 0
			halfdays_count = 0
			halfdays_absent_count = 0
			for doc_name in present_doc_name_list:
				attendance_date = doc_name['attendance_date']
				if not check_holiday(attendance_date, i.employee):
					present_doc_name = frappe.db.get_value(
						"Attendance", 
						filters={
							'employee': i.employee,
							'status': 'Present',  
							'attendance_date': attendance_date,
							'docstatus': ["!=", 2], 
						},
						fieldname=['attendance_date', 'name']
					)
					if present_doc_name:
						present_days_count_ += 1
					halfday_doc_name = frappe.db.get_value(
						"Attendance", 
						filters={
							'employee': i.employee,
							'status': 'Half Day',  
							'attendance_date': attendance_date,
							'docstatus': ["!=", 2], 
							'total_working_hours':[">=",4]
						},
						fieldname=['attendance_date', 'name']
					)
					if halfday_doc_name:
						halfdays_count += 1
					halfday_absent_doc_name = frappe.db.get_value(
					"Attendance", 
					filters={
						'employee': i.employee,
						'status': 'Half Day',  
						'attendance_date': attendance_date,
						'docstatus': ["!=", 2], 
						'total_working_hours':["<",4]
					},
					fieldname=['attendance_date', 'name']
					) 
					if halfday_absent_doc_name:
						halfdays_absent_count += 1

			# totalpresent = present _days_count_ + (halfdays_count / 2)
			totalpresent= i.total_present_days_of_employee
			# cof = frappe.db.sql("""
			# select SUM(total_leave_days) as total_leave_days
			# from `tabLeave Application`
			# where docstatus = 1 and employee = %(employee)s and status = 'Approved' and leave_type= "Compensatory Off"
			#     and (from_date between %(from_date)s and %(to_date)s
			#         or to_date between %(from_date)s and %(to_date)s
			#         or (from_date < %(from_date)s and to_date > %(to_date)s))
			# """, {
			#     "from_date": from_date,
			#     "to_date": to_date,
			#     "employee":i.employee,
			# },
			# as_dict = True)
			# if cof and cof[0].get('total_leave_days') is not None:
			#     tot_coff = cof[0].get('total_leave_days') 
			# else:
			#     tot_coff = 0
			# cl = frappe.db.sql("""
			# select SUM(total_leave_days) as total_leave_days
			# from `tabLeave Application`
			# where docstatus = 1 and employee = %(employee)s and status = 'Approved' and leave_type= "Casual Leave"
			#     and (from_date between %(from_date)s and %(to_date)s
			#         or to_date between %(from_date)s and %(to_date)s
			#         or (from_date < %(from_date)s and to_date > %(to_date)s))
			# """, {
			#     "from_date": from_date,
			#     "to_date": to_date,
			#     "employee":i.employee,
			# },
			# as_dict = True)
			# if cl and cl[0].get('total_leave_days') is not None:
			#     tot_cll = cl[0].get('total_leave_days') 
			# else:
			#     tot_cll = 0
			# sll = frappe.db.sql("""
			# select SUM(total_leave_days) as total_leave_days
			# from `tabLeave Application`
			# where docstatus = 1 and employee = %(employee)s and status = 'Approved' and leave_type= "Sick Leave"
			#     and (from_date between %(from_date)s and %(to_date)s
			#         or to_date between %(from_date)s and %(to_date)s
			#         or (from_date < %(from_date)s and to_date > %(to_date)s))
			# """, {
			#     "from_date": from_date,
			#     "to_date": to_date,
			#     "employee":i.employee,
			# },
			# as_dict = True)
			# if sll and sll[0].get('total_leave_days') is not None:
			#     tot_sll = sll[0].get('total_leave_days') 
			# else:
			#     tot_sll = 0
			# el = frappe.db.sql("""
			#     select SUM(total_leave_days) as total_leave_days
			#     from `tabLeave Application`
			#     where docstatus = 1 and employee = %(employee)s and status = 'Approved' and leave_type= "Earned Leave"
			#         and (from_date between %(from_date)s and %(to_date)s
			#             or to_date between %(from_date)s and %(to_date)s
			#             or (from_date < %(from_date)s and to_date > %(to_date)s))
			#     """, {
			#         "from_date": from_date,
			#         "to_date": to_date,
			#         "employee":i.employee,
			#     },
			#     as_dict = True)
			# if el and el[0].get('total_leave_days') is not None:
			#     tot_earned = el[0].get('total_leave_days') 
			# else:
			#     tot_earned = 0
			# leavepay = frappe.db.sql("""
			# select SUM(total_leave_days) as total_leave_days
			# from `tabLeave Application`
			# where docstatus = 1 and employee = %(employee)s and status = 'Approved' and leave_type= "Leave Without Pay"
			#     and (from_date between %(from_date)s and %(to_date)s
			#         or to_date between %(from_date)s and %(to_date)s
			#         or (from_date < %(from_date)s and to_date > %(to_date)s))
			# """, {
			#     "from_date": from_date,
			#     "to_date": to_date,
			#     "employee":i.employee,
			# },
			# as_dict = True)
			# if leavepay and leavepay[0].get('total_leave_days') is not None:
			#     tot_leavepay = leavepay[0].get('total_leave_days') 
			# else:
			#     tot_leavepay = 0
			tot_coff = get_leave_count(i.employee, "Compensatory Off", from_date, to_date)
			tot_sll = get_leave_count(i.employee, "Sick Leave", from_date, to_date)
			tot_earned = get_leave_count(i.employee, "Earned Leave", from_date, to_date)
			tot_cll = get_leave_count(i.employee, "Casual Leave", from_date, to_date)
			tot_leavepay = get_leave_count(i.employee, "Leave Without Pay", from_date, to_date)
			# tot_leavepay = (halfdays_absent_count / 2) + lop_count_leave
			for j in emp:
				basic=round(frappe.get_value("Salary Detail",{'parent':i.name,'salary_component':'Basic'},['amount']) or 0.0 ) 
				house=round(frappe.get_value("Salary Detail",{'parent':i.name,'salary_component':'House Rent Allowance'},['amount']) or 0.0) 
				special=round(frappe.get_value("Salary Detail",{'parent':i.name,'salary_component':'Special Allowance'},['amount']) or 0.0) 
				heat=round(frappe.get_value("Salary Detail",{'parent':i.name,'salary_component':'Heat Allowance'},['amount']) or 0.0 ) 
				Tl=round(frappe.get_value("Salary Detail",{'parent':i.name,'salary_component':'TL Allowance'},['amount']) or 0.0 ) 
				gratuity=round(frappe.get_value("Salary Detail",{'parent':i.name,'salary_component':'Gratuity'},['amount']) or 0.0 ) 
				other_allowance=round(frappe.get_value("Salary Detail",{'parent':i.name,'salary_component':'Arear Allowance'},['amount']) or 0.0 )
				travel=round(frappe.get_value("Salary Detail",{'parent':i.name,'salary_component':'Travel Allowance'},['amount']) or 0.0) 
				bonus=round(frappe.get_value("Salary Detail",{'parent':i.name,'salary_component':'Bonus'},['amount']) or 0.0) 
				arrearallow=round(frappe.get_value("Salary Detail",{'parent':i.name,'salary_component':'Arear Allowance'},['amount']) or 0.0) 
				overtimeamount=round(frappe.get_value("Salary Detail",{'parent':i.name,'salary_component':'Overtime'},['amount']) or 0.0) 
				attendance=round(frappe.get_value("Salary Detail",{'parent':i.name,'salary_component':'Attendance Allowance'},['amount']) or 0.0) 
				pfound=round(frappe.get_value("Salary Detail",{'parent':i.name,'salary_component':'Provident Fund'},['amount']) or 0.0) 
				esi=round(frappe.get_value("Salary Detail",{'parent':i.name,'salary_component':'Employer State Insurance'},['amount']) or 0.0) 
				protax=round(frappe.get_value("Salary Detail",{'parent':i.name,'salary_component':'Professional Tax'},['amount']) or 0.0) 
				labfund=round(frappe.get_value("Salary Detail",{'parent':i.name,'salary_component':'Labour Welfare Fund'},['amount']) or 0.0) 
				loan=round(frappe.get_value("Salary Detail",{'parent':i.name,'salary_component':'loan'},['amount']) or 0.0) 
				tds=round(frappe.get_value("Salary Detail",{'parent':i.name,'salary_component':'TDS Deduction'},['amount']) or 0.0) 
				others=round(frappe.get_value("Salary Detail",{'parent':i.name,'salary_component':'Others'},['amount']) or 0.0) 
				elencash=round(frappe.get_value("Salary Detail",{'parent':i.name,'salary_component':'EL Encashment'},['amount']) or 0.0) 
				# ss=check_holiday_ss(from_date,to_date,i.employee)
				# nh_data=get_holiday_count(i.employee,from_date,to_date)
				joining_date, relieving_date = frappe.get_cached_value(
					"Employee", i.employee, ["date_of_joining", "relieving_date"]
				)
				nh_data = get_holiday_count(i.employee, from_date, to_date)
				if relieving_date and (getdate(from_date) <= relieving_date < getdate(to_date)):
					frappe.log_error(title="Relieving Date Type", message=f"Type: {type(relieving_date)}, Value: {relieving_date}")
					nh_data = get_holiday_count(i.employee, from_date, relieving_date)
				if joining_date and (getdate(from_date) < joining_date <= getdate(to_date)):
					nh_data = get_holiday_count(i.employee, joining_date, to_date)  
				nh_days = nh_data['total_other_holiday']
				ss_days=nh_data['total_holiday']
				# nh_days=int(nh)
				tot_fixed=j.basic+j.house_rent_allowance+j.special_allowance
				tot_gross=basic+house+special
				overtime=i.ot_hours
				# elencash=0
				canteen=0
				licform=0
			
				fromdate = datetime.strptime(str(j.date_of_birth), '%Y-%m-%d')
				dob = fromdate.strftime('%d-%m-%Y')
				fromdate_ = datetime.strptime(str(j.date_of_joining), '%Y-%m-%d')
				doj = fromdate_.strftime('%d-%m-%Y')
				row=([count, 
				j.employee,
				i.employee_name,
				dob,
				j.grade, 
				doj,  
				i.department, 
				i.designation,
				j.aadhar_no,
				j.pan_no,
				j.uan_no,
				j.esic_no,
				i.bank_account_no, 
				i.bank_type or "", 
				j.employee_category,
				i.total_working_days,
				j.basic,
				j.house_rent_allowance,
				j.special_allowance,
				tot_fixed,
				totalpresent,ss_days,tot_coff,tot_cll,tot_sll,tot_earned,nh_days,tot_leavepay,i['absent_days'],i.payment_days,
				basic,house,special,
				tot_gross,
				attendance,heat,Tl,elencash,gratuity,travel,bonus,overtime,overtimeamount,other_allowance,
				round(i.gross_pay,0), pfound,round(esi,0), protax,labfund,loan,others,tds,
				round(i.total_deduction),
				round(i.net_pay)
				])    
				total_workingdays+=int(i.total_working_days)	
				fixed_basic+=float(j.basic)
				fixed_house+=float(j.house_rent_allowance)
				fixed_special+=float(j.special_allowance)
				fixed_total+=tot_fixed
				total_presentday+=float(totalpresent)
				total_ss+=ss_days
				total_tot_coff+=tot_coff
				total_tot_cll+=tot_cll
				total_tot_sll+=tot_sll
				total_tot_earned+=tot_earned
				total_nh+=nh_days
				total_tot_leavepay+=tot_leavepay
				absent_total += i['absent_days']
				total_payment_day+=i.payment_days
				gross_basic+=int(basic)
				gross_house+=int(house)
				gross_special+=int(special)
				gross_total+=tot_gross
				variable_attendance+=attendance
				variable_heat+=heat
				variable_tl+=Tl
				variable_other+=other_allowance
				variable_gratuity+=gratuity
				variable_travel+=travel
				variable_bonus+=bonus
				variable_overtime+=overtime
				variable_overtimeamount+=overtimeamount
				variable_arrearallow+=arrearallow
				variable_elencash+=int(elencash)
				total_grosspay+=round(i.gross_pay)
				deduction_pfound+=int(pfound)
				deduction_esi+=round(esi,0)
				deduction_protax+=int(protax)
				deduction_canteen+=int(canteen)
				deduction_licform+=int(licform)
				deduction_labfund+=int(labfund)
				deduction_labfund+=int(loan)
				deduction_others+=int(others)
				deduction_tds+=int(tds)		
				deduction_total_deduction+=round(i.total_deduction)
				total_netpay+=round(i.net_pay)

			data.append(row)

	row1=[ 'TOTAL','','','','','','','','','','','','','','',total_workingdays,fixed_basic,fixed_house,fixed_special,fixed_total,total_presentday,
		total_ss, total_tot_coff, total_tot_cll, total_tot_sll,total_tot_earned,total_nh, total_tot_leavepay,absent_total,
		total_payment_day, gross_basic,gross_house,gross_special,gross_total,variable_attendance, variable_heat,variable_tl,variable_elencash,variable_gratuity,variable_travel,variable_bonus,
		variable_overtime, variable_overtimeamount, variable_other, total_grosspay,deduction_pfound,
		round(deduction_esi,0), deduction_protax, deduction_labfund,deduction_loan,deduction_others,deduction_tds,
		deduction_total_deduction, total_netpay]
	data.append(row1)
	return data                

@frappe.whitelist()
def get_dates(from_date,to_date):
	no_of_days = date_diff(add_days(to_date, 1), from_date)
	dates = [add_days(from_date, i) for i in range(0, no_of_days)]
	return dates

@frappe.whitelist()
def check_holiday_ss(from_date, to_date, emp):
	dates = get_dates(from_date, to_date)
	tot = 0
	for date in dates:
		holiday_list = frappe.db.get_value('Employee', {'name': emp}, 'holiday_list')
		holiday = frappe.db.sql("""
			SELECT tabHoliday.holiday_date, tabHoliday.weekly_off
			FROM `tabHoliday List` AS holiday_list
			LEFT JOIN `tabHoliday` ON tabHoliday.parent = holiday_list.name
			WHERE holiday_list.name = %s AND holiday_date = %s
		""", (holiday_list, date), as_dict=True)
		
		doj = frappe.db.get_value("Employee", {'name': emp}, "date_of_joining")
		
		if holiday:
			if doj < holiday[0].holiday_date:
				if holiday[0].weekly_off == 1 or holiday[0].holiday_type == 'PH':
					tot += 1
	return tot

@frappe.whitelist()
def check_holiday_nh(from_date, to_date, emp):
	dates = get_dates(from_date, to_date)
	tot = 0
	for date in dates:
		holiday_list = frappe.db.get_value('Employee', {'name': emp}, 'holiday_list')
		holiday = frappe.db.sql("""
			SELECT tabHoliday.holiday_date, tabHoliday.weekly_off
			FROM `tabHoliday List` AS holiday_list
			LEFT JOIN `tabHoliday` ON tabHoliday.parent = holiday_list.name
			WHERE holiday_list.name = %s AND holiday_date = %s
		""", (holiday_list, date), as_dict=True)
		
		doj = frappe.db.get_value("Employee", {'name': emp}, "date_of_joining")
		
		if holiday:
			if doj < holiday[0].holiday_date:
				if holiday[0].holiday_type == 'NH' or holiday[0].holiday_type == 'FH':
					tot += 1
	return tot



def check_holiday(date, emp):
	# Get the holiday list for the employee
	holiday_list = frappe.db.get_value('Employee', emp, 'holiday_list')
	
	# Parameterized query to prevent SQL injection and ensure proper handling of input values
	query = """
		SELECT `tabHoliday`.holiday_date, `tabHoliday`.weekly_off
		FROM `tabHoliday List`
		LEFT JOIN `tabHoliday` ON `tabHoliday`.parent = `tabHoliday List`.name
		WHERE `tabHoliday List`.name = %s AND holiday_date = %s
	"""
	holiday = frappe.db.sql(query, (holiday_list, date), as_dict=True)
	if holiday:
		if holiday[0].weekly_off == 1:
			return "WW"
		else:
			return "HH"
	return None


def get_holiday_count(employee, from_date, to_date):
	holiday_list = frappe.db.get_value('Employee', employee, 'holiday_list')
	total_weekoff = 0
	total_public_holiday = 0
	total_other_holiday = 0

	if holiday_list:
		holiday_records = frappe.db.sql("""
			SELECT holiday_date, weekly_off, holiday_type
			FROM `tabHoliday`
			WHERE parent = %s AND holiday_date BETWEEN %s AND %s
		""", (holiday_list, from_date, to_date), as_dict=True)

		for record in holiday_records:
			if record['weekly_off'] == 1:
				total_weekoff += 1
			elif record['holiday_type'] == "PH":
				total_public_holiday += 1
			else:
				total_other_holiday += 1

	return {
		'total_weekoff': total_weekoff,
		'total_public_holiday': total_public_holiday,
		'total_other_holiday': total_other_holiday,
		'total_holiday': total_weekoff + total_public_holiday
	}
	  